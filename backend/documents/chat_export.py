# -*- coding: utf-8 -*-
"""把已保存的助手回答导出为 Markdown / Word。

对应实施计划 3.2：

★ 不再次调用模型，只做格式转换；
★ 保留标题、段落、列表、表格与已有的真实引用（[S1] 之类原样保留）；
★ 来源只取该条消息里真实存在的引用，不补造。
"""
from __future__ import annotations

import io
import re
from datetime import datetime

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
_ORDERED = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
_TABLE_RULE = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")
_QUOTE = re.compile(r"^>\s?(.*)$")
_FENCE = re.compile(r"^\s*```")
_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*\n]+?\*(?!\*))")


def markdown_bytes(text: str) -> bytes:
    """Markdown 原样保存为 UTF-8（不加 BOM，便于通用编辑器识别）。"""
    return (text or "").encode("utf-8")


def _add_inline(paragraph, text: str) -> None:
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def _cells(line: str) -> list[str]:
    body = line.strip()
    if body.startswith("|"):
        body = body[1:]
    if body.endswith("|"):
        body = body[:-1]
    return [cell.strip() for cell in body.split("|")]


def docx_bytes(text: str, *, title: str | None = None,
               sources: list[dict] | None = None) -> bytes:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    # 中文字体要单独设 eastAsia，否则 python-docx 只改西文
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    if title:
        document.add_heading(title, level=1)

    lines = (text or "").replace("\r\n", "\n").split("\n")
    index = 0
    in_fence = False
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if _FENCE.match(line):
            in_fence = not in_fence
            index += 1
            continue
        if in_fence:
            run = document.add_paragraph().add_run(line)
            run.font.name = "Consolas"
            index += 1
            continue
        if not stripped:
            index += 1
            continue

        heading = _HEADING.match(stripped)
        if heading:
            document.add_heading(heading.group(2).strip(), level=min(len(heading.group(1)), 6))
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and _TABLE_RULE.match(lines[index + 1]):
            header = _cells(stripped)
            index += 2
            body = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                body.append(_cells(lines[index]))
                index += 1
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            for position, cell in enumerate(header):
                table.rows[0].cells[position].text = cell
            for row in body:
                cells = table.add_row().cells
                for position in range(len(header)):
                    cells[position].text = row[position] if position < len(row) else ""
            document.add_paragraph()
            continue

        quote = _QUOTE.match(stripped)
        if quote:
            paragraph = document.add_paragraph(style="Intense Quote")
            _add_inline(paragraph, quote.group(1))
            index += 1
            continue

        bullet = _BULLET.match(line)
        if bullet:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_inline(paragraph, bullet.group(2))
            index += 1
            continue

        ordered = _ORDERED.match(line)
        if ordered:
            paragraph = document.add_paragraph(style="List Number")
            _add_inline(paragraph, ordered.group(3))
            index += 1
            continue

        if re.fullmatch(r"[-*_]{3,}", stripped):
            document.add_paragraph("─" * 20)
            index += 1
            continue

        paragraph = document.add_paragraph()
        _add_inline(paragraph, stripped)
        index += 1

    if sources:
        document.add_heading("资料来源", level=2)
        for number, source in enumerate(sources, 1):
            name = str(source.get("title") or source.get("source") or "").strip() or "未命名资料"
            suffix = ("（%s）" % source["url"]) if source.get("url") else ""
            document.add_paragraph("%d. %s%s" % (number, name, suffix), style="List Number")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def filename_for(message_id: int, fmt: str, title: str | None = None) -> str:
    stem = (title or "回答").strip().replace("/", " ").replace("\\", " ")[:40] or "回答"
    stamp = datetime.now().strftime("%Y%m%d")
    return "%s_%s_%s.%s" % (stem, message_id, stamp, fmt)
