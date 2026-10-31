# -*- coding: utf-8 -*-
"""会话内表格文件：上传落盘、归属校验、读取与结果另存。

设计边界（对应实施计划 3.3）：

★ 文件实体放 `data/chat_files/<conversation_id>/`，数据库只存元数据；
  **任何下载都必须经过 `get()` 校验 user_id / conversation_id**，不允许把路径当下载地址。
★ 上传文件永不改写；工具结果写成新的 file_id（role='result'）。
★ 表格读取只向模型暴露列名、行数与少量样本，完整数据留在服务端。
★ CSV 先按 UTF-8/UTF-8 BOM 解码，失败再试 GB18030；都失败则明确报错。
★ `.xlsx` 的公式与缓存值分开看：公式没有缓存值时如实说明无法计算，不假装算过。
"""
from __future__ import annotations

import csv
import io
import uuid
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from core.paths import BACKEND_ROOT
from database.db import ChatFile

STORAGE_ROOT = BACKEND_ROOT / "data" / "chat_files"
ALLOWED_SUFFIXES = {".xlsx", ".csv"}
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
#: 只给模型看这么多行样本，避免整表进上下文。
SAMPLE_ROWS = 5


class FileNotAvailable(LookupError):
    """文件不存在，或不属于当前用户/会话。统一按 404 处理，不区分两者。"""


class UnsupportedSheet(ValueError):
    """文件能读但无法可靠计算（例如公式没有缓存值）。"""


@dataclass(frozen=True)
class TableInfo:
    sheets: list[str]
    sheet: str
    columns: list[str]
    rows: int
    sample: list[dict]
    encoding: str | None
    formula_cells: int
    uncached_formulas: int


def _folder(conversation_id: int) -> Path:
    folder = STORAGE_ROOT / str(int(conversation_id))
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def drop_conversation_folder(conversation_id: int) -> None:
    """会话被删除后清掉它自己的文件目录（只动这个会话目录，失败不抛出）。"""
    folder = STORAGE_ROOT / str(int(conversation_id))
    if not folder.is_dir():
        return
    for item in folder.iterdir():
        try:
            item.unlink()
        except OSError:
            return
    try:
        folder.rmdir()
    except OSError:
        pass


def _safe_name(filename: str, suffix: str) -> str:
    """保留原文件名用于展示，但去掉路径分隔符与控制字符。"""
    name = Path(str(filename or "上传文件")).name
    name = "".join(ch for ch in name if ch.isprintable() and ch not in '\\/:*?"<>|').strip()
    if not name:
        name = "上传文件" + suffix
    if not name.lower().endswith(suffix):
        name += suffix
    return name[:120]


def save_bytes(db, *, user_id: int, conversation_id: int, filename: str,
               data: bytes, role: str = "upload") -> ChatFile:
    suffix = Path(filename or "").suffix.lower()
    allowed = ALLOWED_SUFFIXES | ({'.docx', '.md'} if role == 'result' else set())
    if suffix not in allowed:
        raise UnsupportedSheet("文件格式不受支持")
    if not data:
        raise UnsupportedSheet("文件是空的")
    if len(data) > MAX_UPLOAD_BYTES:
        raise UnsupportedSheet("文件超过 20 MiB，请先拆分成更小的表格")
    if suffix == ".csv":
        _decode_csv(data)  # 先验证能解码，避免存进去之后才发现读不了
    file_id = uuid.uuid4().hex
    target = _folder(conversation_id) / (file_id + suffix)
    target.write_bytes(data)
    row = ChatFile(id=file_id, conversation_id=int(conversation_id), user_id=int(user_id),
                   role=role, filename=_safe_name(filename, suffix),
                   path=str(target), size=len(data), created_at=datetime.utcnow())
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def save_result(db, *, user_id: int, conversation_id: int, filename: str,
                data: bytes) -> ChatFile:
    return save_bytes(db, user_id=user_id, conversation_id=conversation_id,
                      filename=filename, data=data, role="result")


def get(db, file_id: str, *, user_id: int, conversation_id: int | None = None) -> ChatFile:
    row = db.query(ChatFile).filter(ChatFile.id == str(file_id or "")).first()
    if row is None or int(row.user_id) != int(user_id):
        raise FileNotAvailable("文件不存在或无权访问")
    if conversation_id is not None and int(row.conversation_id) != int(conversation_id):
        raise FileNotAvailable("文件不属于当前会话")
    path = Path(row.path)
    if not path.is_file():
        raise FileNotAvailable("文件已被清理，请重新上传")
    return row


def public(row: ChatFile) -> dict:
    return {"file_id": row.id, "filename": row.filename, "size": int(row.size or 0),
            "role": row.role}


def _decode_csv(data: bytes) -> tuple[str, str]:
    # 先看 BOM 再决定编码，否则纯 UTF-8 文件也会被报成 utf-8-sig，误导用户。
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig"), "utf-8-sig"
    for encoding in ("utf-8", "gb18030"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise UnsupportedSheet("CSV 无法按 UTF-8 或 GB18030 解码，请另存为 UTF-8 后重试")


def _formula_stats(path: Path, sheet: str | None) -> tuple[int, int]:
    """统计公式单元格数量，以及"有公式但没有缓存值"的数量。"""
    if path.suffix.lower() != ".xlsx":
        return 0, 0
    import openpyxl
    try:
        formulas = openpyxl.load_workbook(path, data_only=False, read_only=True)
        cached = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return 0, 0
    names = [sheet] if sheet and sheet in formulas.sheetnames else formulas.sheetnames[:1]
    total = missing = 0
    for name in names:
        try:
            fs, cs = formulas[name], cached[name]
        except KeyError:
            continue
        for f_row, c_row in zip(fs.iter_rows(values_only=True), cs.iter_rows(values_only=True)):
            for f_value, c_value in zip(f_row, c_row):
                if isinstance(f_value, str) and f_value.startswith("="):
                    total += 1
                    if c_value is None:
                        missing += 1
    return total, missing


def _frame(path: Path, sheet: str | None):
    import pandas as pd
    if path.suffix.lower() == ".csv":
        text, encoding = _decode_csv(path.read_bytes())
        return pd.read_csv(io.StringIO(text)), encoding
    book = pd.ExcelFile(path)
    name = sheet if sheet in book.sheet_names else book.sheet_names[0]
    return pd.read_excel(book, sheet_name=name), None


def describe(row: ChatFile, sheet: str | None = None) -> TableInfo:
    import pandas as pd
    path = Path(row.path)
    formula_total, formula_missing = _formula_stats(path, sheet)
    if path.suffix.lower() == ".xlsx":
        names = [str(name) for name in pd.ExcelFile(path).sheet_names]
    else:
        names = ["(csv)"]
    active = sheet if sheet in names else names[0]
    frame, encoding = _frame(path, active)
    columns = [str(column) for column in frame.columns]
    sample_rows = [{str(key): _plain(value) for key, value in record.items()}
                   for record in frame.head(SAMPLE_ROWS).to_dict(orient="records")]
    return TableInfo(sheets=names, sheet=active, columns=columns, rows=int(len(frame)),
                     sample=sample_rows, encoding=encoding, formula_cells=formula_total,
                     uncached_formulas=formula_missing)


def _plain(value):
    if value is None:
        return None
    if isinstance(value, float) and value != value:  # NaN
        return None
    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            return str(value)
    return value
