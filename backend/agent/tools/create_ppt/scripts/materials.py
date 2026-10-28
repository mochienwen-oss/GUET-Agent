"""Task-local source documents and data; never writes the public corpus."""
import csv
import io
import json
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from pydantic import BaseModel, ConfigDict, Field, model_validator

MAX_UPLOAD_BYTES = 20 * 1024 * 1024


class Material(BaseModel):
    model_config = ConfigDict(extra="forbid")
    material_id: str = Field(default_factory=lambda: uuid4().hex[:12], pattern=r"^[a-f0-9]{12}$")
    title: str = Field(min_length=1, max_length=300)
    text: str = Field(default="", max_length=200000)
    source_type: str = "user_text"
    locator: str = "用户提供的文本"
    url: str | None = None
    date: str | None = None
    scope: str = ""
    columns: list[str] = Field(default_factory=list)
    rows: list[list[str | int | float | None]] = Field(default_factory=list)
    units: dict[str, str] = Field(default_factory=dict)
    asset_file: str | None = Field(default=None, pattern=r"^[a-f0-9]{12}\.png$")

    @model_validator(mode="after")
    def usable_content(self):
        if not self.title.strip() or not (self.text.strip() or self.columns):
            raise ValueError("材料标题和内容不能为空")
        if self.columns:
            if len(set(self.columns)) != len(self.columns) or any(not c.strip() for c in self.columns):
                raise ValueError("数据表字段须非空且唯一")
            if not self.rows or len(self.rows) > 10000 or len(self.columns) > 100:
                raise ValueError("数据表需有1—10000行、至多100列")
            if any(len(row) != len(self.columns) for row in self.rows):
                raise ValueError("数据行与字段数不一致")
        return self

    def evidence(self):
        text = self.text
        if self.columns:
            text = json.dumps({"columns": self.columns, "rows": self.rows, "units": self.units}, ensure_ascii=False)
        return {"evidence_id": self.material_id, "title": self.title, "text": text,
                "heading_path": [self.locator], "source_type": self.source_type,
                "url": self.url, "date": self.date, "scope": self.scope,
                "material_id": self.material_id}


def _table(title, values, locator):
    rows = []
    for row in values:
        if any(v is not None and v != "" for v in row):
            rows.append([str(v) if isinstance(v, bool) else v.isoformat() if hasattr(v, "isoformat") else v for v in row])
        if len(rows) > 10001:
            raise ValueError("数据表超过10000行，请筛选所需范围")
    if len(rows) < 2:
        raise ValueError("表格须包含字段行和数据行")
    columns = [str(c).strip() if c is not None else "" for c in rows[0]]
    return Material(title=title, source_type="data", locator=locator,
                    columns=columns, rows=rows[1:])


def read_materials(filename: str, content: bytes, *, asset_dir: Path | None = None) -> list[Material]:
    """Decode an uploaded byte stream; client paths never become file paths."""
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("文件为空或超过20MiB")
    name = Path(filename.replace("\\", "/")).name
    suffix = Path(name).suffix.lower()
    if suffix in {'.png', '.jpg', '.jpeg', '.webp'}:
        if asset_dir is None:
            raise ValueError('图片须通过当前任务的材料上传入口添加')
        from PIL import Image, UnidentifiedImageError
        try:
            with Image.open(io.BytesIO(content)) as image:
                if image.width * image.height > 25000000:
                    raise ValueError('图片超过2500万像素，请缩小后上传')
                width, height = image.size
                normalized = image.convert('RGBA' if 'A' in image.getbands() or 'transparency' in image.info else 'RGB')
        except (UnidentifiedImageError, Image.DecompressionBombError):
            raise ValueError('图片格式损坏或尺寸过大') from None
        material = Material(title=name, source_type='user_image', locator='用户上传图片',
                            text='用户提供的图片素材，尺寸%d×%d。可用于页面配图；具体含义以用户说明为准，不据此猜测未识别的内容。' % (width, height))
        asset_dir.mkdir(parents=True, exist_ok=True)
        material.asset_file = material.material_id + '.png'
        normalized.save(asset_dir / material.asset_file, format='PNG')
        return [material]
    if suffix in {".docx", ".xlsx"}:
        from zipfile import is_zipfile
        if not is_zipfile(io.BytesIO(content)):
            raise ValueError("文档格式损坏，请重新导出文件")
        with ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("文档解压后超过100MiB，请缩小材料范围")
    if suffix in {".txt", ".md", ".csv"}:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("文本或CSV请保存为UTF-8编码") from None
        if suffix == ".csv":
            return [_table(name, csv.reader(io.StringIO(text)), "CSV字段与数据行")]
        return [Material(title=name, text=text, source_type="user_document", locator="全文")]
    if suffix == ".docx":
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        doc = Document(io.BytesIO(content))
        blocks = []
        for node in doc.element.body.iterchildren():
            if node.tag.endswith("}p"):
                blocks.append(Paragraph(node, doc).text)
            elif node.tag.endswith("}tbl"):
                blocks.append("\n".join(" | ".join(c.text for c in row.cells) for row in Table(node, doc).rows))
        return [Material(title=name, text="\n\n".join(blocks),
                         source_type="user_document", locator="正文段落与文档表格")]
    if suffix == ".pdf":
        import fitz
        with fitz.open(stream=content, filetype="pdf") as doc:
            text = "\n\n".join("[第%d页]\n%s" % (i + 1, page.get_text())
                                 for i, page in enumerate(doc))
            if not any(page.get_text().strip() for page in doc):
                raise ValueError("PDF没有可提取的文字，请提供文字版材料")
        return [Material(title=name, text=text, source_type="user_document", locator="正文内标有原始页码")]
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        result = []
        try:
            for sheet in book:
                from itertools import islice
                values = list(islice(sheet.iter_rows(values_only=True), 10002))
                if len(values) > 10001:
                    raise ValueError("工作表超过10000数据行，请先筛选或移除空白尾行；未截断数据")
                if not any(any(v is not None for v in row) for row in values):
                    continue
                if any(isinstance(v, str) and v.startswith("=") for row in values for v in row):
                    raise ValueError("首版读取数值表；含公式的工作表请先导出计算结果为CSV")
                result.append(_table(name + " · " + sheet.title, values, "工作表：" + sheet.title))
        finally:
            book.close()
        if not result:
            raise ValueError("工作簿没有可用数据")
        return result
    raise ValueError("支持TXT、Markdown、DOCX、文字PDF、CSV、XLSX及PNG/JPEG/WebP图片")
