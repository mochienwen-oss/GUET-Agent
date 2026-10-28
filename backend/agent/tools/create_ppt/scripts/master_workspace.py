"""Task-local access to the pinned PPT Master package and its real tools."""
import base64
import csv
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET

from agent.core.request_control import check_cancelled, remaining_seconds
from agent.skills.ppt_master import skill_dir

#: 由 skill 模块解析，不再写死仓库路径；升级/搬迁只需要改 manifest.json。
SKILL_DIR = skill_dir('ppt_master')
SCRIPTS = {
    'attribution_guard.py',
    'project_manager.py', 'text_measure.py', 'preset_shape_svg.py',
    'shape_boolean_svg.py', 'svg_position_calculator.py', 'stamp_native_fallbacks.py',
    'svg_quality_checker.py', 'total_md_split.py', 'finalize_svg.py',
    'svg_to_pptx.py', 'semantic_table.py', 'native_payloads.py', 'apply_template.py',
    'icon_sync.py', 'visualization_recall.py', 'analyze_images.py',
    'animation_config.py', 'update_spec.py', 'workflow_log.py',
}


def inside(root, path):
    root = Path(root).resolve()
    target = (root / path).resolve()
    if not target.is_relative_to(root) or any(p.startswith('.env') for p in target.parts):
        raise ValueError('路径不属于当前任务或允许的Skill资源')
    return target


def snapshot_bytes(snapshot, name):
    """修订基准以 base64 存在 state 里（可 JSON 序列化），取用时还原成字节。"""
    value = (snapshot or {}).get(name)
    if value is None:
        return b''
    if isinstance(value, bytes):
        return value
    try:
        return base64.b64decode(value)
    except (ValueError, TypeError):
        return b''


def slide_files(folder):
    scripts = str(SKILL_DIR / 'scripts')
    if scripts not in sys.path:
        sys.path.append(scripts)
    from slide_roster import discover_slide_svgs
    return discover_slide_svgs(Path(folder))


def page_content(slide):
    assets = []
    def collect(shapes):
        for shape in shapes:
            if hasattr(shape, 'image'):
                assets.append(shape.image.blob)
            if shape.has_chart:
                assets.append(shape.chart.part.blob)
            if hasattr(shape, 'shapes'):
                collect(shape.shapes)
    collect(slide.shapes)
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ''
    return slide._element.xml, assets, notes


def command(script, arguments, project, stdin=None):
    """No shell, interpreter snippets, arbitrary programs, or inherited API keys."""
    check_cancelled()
    environment = {k: v for k, v in os.environ.items()
                   if k.upper() in {'PATH', 'SYSTEMROOT', 'WINDIR', 'TEMP', 'TMP', 'LANG', 'LD_LIBRARY_PATH'}}
    environment.update(PYTHONIOENCODING='utf-8', PYTHONDONTWRITEBYTECODE='1',
                       PPT_MASTER_PROJECT_PATH=str(project),
                       USERPROFILE=str(Path(project).resolve()))
    input_options = {'input': stdin} if stdin is not None else {'stdin': subprocess.DEVNULL}
    result = subprocess.run([sys.executable, str(SKILL_DIR / 'scripts' / script), *arguments],
                            cwd=project, env=environment, capture_output=True, text=True,
                            encoding='utf-8', errors='replace', timeout=remaining_seconds(120), **input_options)
    return {'exit_code': result.returncode, 'stdout': result.stdout, 'stderr': result.stderr,
            'output': result.stdout + result.stderr}


class MasterWorkspace:
    def __init__(self, folder):
        self.folder = Path(folder).resolve()

    def initialize(self, state):
        self.folder.parent.mkdir(parents=True, exist_ok=True)
        guard = command('attribution_guard.py', [], self.folder.parent)
        if guard['exit_code']:
            raise ValueError('PPT Master入口完整性检查未通过，已停止，未绕过检查')
        if not self.folder.exists():
            result = command('project_manager.py', ['init', self.folder.name,
                             '--dir', str(self.folder.parent)], self.folder.parent)
            if result['exit_code']:
                raise ValueError('PPT Master项目初始化失败：' + result['output'][-1200:])
        self.import_materials(state)

    def import_materials(self, state):
        sources = self.folder / 'sources'
        sources.mkdir(exist_ok=True)
        (sources / 'brief.json').write_text(json.dumps(state['brief'], ensure_ascii=False, indent=2), encoding='utf-8')
        (sources / 'evidence.json').write_text(json.dumps(state.get('evidence', {}), ensure_ascii=False, indent=2), encoding='utf-8')
        (sources / 'analyses.json').write_text(json.dumps(state.get('analyses', []), ensure_ascii=False, indent=2), encoding='utf-8')
        for material in state.get('materials') or []:
            stem = material['material_id']
            asset_note = ''
            if material.get('asset_file'):
                origin = inside(Path(self.folder).parents[1] / 'uploads', material['asset_file'])
                image = inside(self.folder / 'images', material['asset_file'])
                image.parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(origin, image)
                asset_note = '\n\n任务图片路径：images/' + material['asset_file']
            (sources / (stem + '.md')).write_text(
                '# ' + material['title'] + '\n\n' + json.dumps({k: material.get(k) for k in
                ('material_id', 'scope', 'date', 'locator', 'units')}, ensure_ascii=False) + '\n\n' +
                material.get('text', '') + asset_note, encoding='utf-8')
            if material.get('columns'):
                stream = io.StringIO(newline='')
                writer = csv.writer(stream)
                writer.writerow(material['columns'])
                writer.writerows(material['rows'])
                (sources / (stem + '.csv')).write_text(stream.getvalue(), encoding='utf-8')

    def read(self, path, *, resource=False, start=0, count=160):
        root = SKILL_DIR if resource else self.folder
        target = inside(root, path)
        if target.suffix.lower() not in {'.md', '.json', '.svg', '.txt', '.csv'}:
            raise ValueError('只读取文字资料；不读取密钥、二进制或程序文件')
        lines = target.read_text(encoding='utf-8-sig').splitlines()
        if start < 0 or not 1 <= count <= 240:
            raise ValueError('读取范围无效')
        if sum(len(line) for line in lines[start:start + count]) > 40000:
            raise ValueError('本次读取超过40,000字符，请选择更小的章节或行范围；未静默截断')
        return {'path': path, 'total_lines': len(lines), 'start': start,
                'text': '\n'.join(lines[start:start + count])}

    def listing(self, path='', *, resource=False):
        root = SKILL_DIR if resource else self.folder
        target = inside(root, path)
        return sorted(p.name + ('/' if p.is_dir() else '') for p in target.iterdir()
                      if not p.name.startswith('.'))[:200]

    def write(self, path, text, *, revision_files=None):
        target = inside(self.folder, path)
        relative = target.relative_to(self.folder)
        if (relative.parts[0] in {'sources', 'exports', 'validation'}
                or target.suffix not in {'.md', '.json', '.svg', '.txt', '.csv'}):
            raise ValueError('只能写设计、页面及备注文件；来源、校验和导出由宿主/上游工具维护')
        if revision_files is not None and relative.parts[0] == 'svg_output' and target.name not in revision_files:
            raise ValueError('不能修改未选页面')
        if target.suffix == '.svg':
            tree = ET.fromstring(text)
            for element in tree.iter():
                for name, value in element.attrib.items():
                    if name.split('}')[-1] == 'href' and value and not value.startswith(('#', 'data:')):
                        if re.match(r'^[a-zA-Z][\w+.-]*:', value):
                            raise ValueError('SVG外部资源须先成为当前任务内的材料')
                        inside(self.folder, target.parent / value)
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(text, encoding='utf-8')
        return {'written': path}

    def run(self, script, arguments, stdin=None):
        if script.replace('\\', '/').split('/')[-1] == 'server.py':
            raise ValueError('当前工作区已承接页面预览与方案确认，不启动独立svg_editor/server.py或确认服务；请在当前工作区查看实际保存的页面')
        if script not in SCRIPTS or not all(isinstance(a, str) for a in arguments):
            raise ValueError('不是允许的PPT Master命令；当前可用脚本：' + '、'.join(sorted(SCRIPTS)))
        if arguments in (['--help'], ['-h']):
            return command(script, arguments, self.folder)
        input_dash = any(value == '--input=-' or value == '--input' and
                         index + 1 < len(arguments) and arguments[index + 1] == '-'
                         for index, value in enumerate(arguments))
        if input_dash and not stdin:
            raise ValueError('当前execute_python接口没有stdin数据通道，不能使用--input -。请先用write_file保存真实JSON文件，再用--input @project/shapes.json调用原版批量脚本')
        if '--stdin' in arguments and not stdin:
            raise ValueError('当前execute_python接口没有stdin数据通道；text_measure的measure/wrap请按原版用法直接传入待测文字')
        if script == 'icon_sync.py':
            if len(arguments) < 2 or (arguments[0] not in {'@project', '@project/'} and inside(self.folder, arguments[0]) != self.folder):
                raise ValueError('图标同步用法：icon_sync.py，arguments=["@project", "library/name", ...]；目标只能是当前项目根目录')
            # These are symbolic bundled ids, not arbitrary paths. The original
            # script still checks its supported libraries, missing names and
            # the one-stylistic-library rule before export.
            return command(script, [str(self.folder), *arguments[1:]], self.folder, stdin)
        if script == 'visualization_recall.py':
            # Its original CLI accepts only semantic tags and catalog ids; it
            # never treats family/key or tag text as paths to open or write.
            return command(script, arguments, self.folder, stdin)
        if script == 'project_manager.py' and (not arguments or arguments[0] not in
                {'validate', 'info', 'page-context', 'page-context-report'}):
            raise ValueError('项目初始化与来源导入由宿主负责')
        expanded = []
        for value in arguments:
            prefix, separator, part = value.partition('=') if value.startswith('--') else ('', '', value)
            if part == '@project' or part.startswith('@project/'):
                part = str(inside(self.folder, part.removeprefix('@project').lstrip('/')))
            elif part.startswith('@skill/'):
                if script != 'apply_template.py' or not expanded or expanded[-1] != '--root':
                    raise ValueError('Skill路径只能作为模板安装的只读输入')
                part = str(inside(SKILL_DIR, part[len('@skill/'):]))
            elif '/' in part or '\\' in part or re.match(r'^[A-Za-z]:', part):
                if script == 'apply_template.py' and expanded and expanded[-1] == '--root':
                    part = str(inside(SKILL_DIR, part))
                else:
                    part = str(inside(self.folder, part))
            expanded.append(prefix + separator + part)
        if script == 'preset_shape_svg.py' and arguments and arguments[0] == 'render-batch':
            input_path = None
            for index, value in enumerate(expanded):
                if value.startswith('--input='):
                    input_path = value.partition('=')[2]
                elif value == '--input' and index + 1 < len(expanded):
                    input_path = expanded[index + 1]
            if input_path is not None and input_path != '-':
                target = inside(self.folder, input_path)
                if target.suffix.lower() != '.json' or not target.is_file():
                    raise ValueError('render-batch的--input须指向已保存的当前任务JSON文件；先write_file写入@project/shapes.json，数组内容由原版校验')
        result = command(script, expanded, self.folder, stdin)
        if script == 'preset_shape_svg.py':
            mode = arguments[0] if arguments else ''
            result.update(page_written=False, output_kind={
                'render': 'svg_fragment', 'render-batch': 'svg_fragments',
                'list': 'preset_catalog', 'describe': 'preset_description'}.get(mode, 'command_output'),
                instruction='本命令只返回stdout内容，不写入页面。render/render-batch返回的是SVG片段，须自行组装完整<svg>并用write_file保存到svg_output/后才有实际页面。')
        return result

    def publish(self, state, session):
        """Check the actual exported deck; never replace it with the legacy renderer."""
        from decimal import Decimal
        from pptx import Presentation
        from .preview import render_previews
        from documents.document_availability import require_ready

        for source in (state.get('evidence') or {}).values():
            if source.get('source_revision') is not None:
                require_ready(source['source_revision'])

        names = [path.name for path in slide_files(self.folder / 'svg_output')]

        def usable_receipt():
            return next((item for item in reversed(session.get('commands') or [])
                         if item['script'] == 'svg_to_pptx.py' and item['exit_code'] == 0
                         and not item.get('stale') and item.get('output')), None)

        def needs_export(receipt):
            """导出凭据与当前页面/报告对不上时，需要宿主自己补跑一次。"""
            if receipt is None:
                return True
            candidate = inside(self.folder, receipt['output'])
            report_file = self.folder / 'validation' / (candidate.stem + '.report.json')
            if not candidate.is_file() or not report_file.is_file():
                return True
            try:
                data = json.loads(report_file.read_text(encoding='utf-8'))
            except (OSError, ValueError):
                return True
            output = data.get('output') or {}
            return (data.get('status') not in {'passed', 'passed-with-warnings'}
                    or Path(str(output.get('path'))).resolve() != candidate.resolve()
                    or output.get('bytes') != candidate.stat().st_size
                    or (data.get('source') or {}).get('svg_slide_count') != len(names))

        receipt = usable_receipt()
        if needs_export(receipt):
            # 自愈：先补齐当前 svg_output 的最终质量报告，再真实导出一次
            check = self.run('svg_quality_checker.py',
                             [str(self.folder), '--canonical-authoring', '--stage', 'final', '--json'])
            if check['exit_code']:
                raise ValueError('自动补导出前的页面检查未通过：' +
                                 ((check.get('output') or '').strip()[-600:] or '原版检查未返回通过结果'))
            before = {path: (path.stat().st_mtime_ns, path.stat().st_size)
                      for path in self.folder.rglob('*.pptx')}
            result = self.run('svg_to_pptx.py', [str(self.folder)])
            for command in session.get('commands') or []:
                if command['script'] == 'svg_to_pptx.py':
                    command['stale'] = True
            record = {'script': 'svg_to_pptx.py', 'arguments': [str(self.folder)],
                      'exit_code': result['exit_code']}
            if result['exit_code'] == 0:
                outputs = [path for path in self.folder.rglob('*.pptx')
                           if before.get(path) != (path.stat().st_mtime_ns, path.stat().st_size)]
                if outputs:
                    record['output'] = str(max(outputs, key=lambda path: path.stat().st_mtime_ns))
            session.setdefault('commands', []).append(record)
            if result['exit_code'] or not record.get('output'):
                raise ValueError('自动重新导出未成功：' +
                                 ((result.get('output') or '').strip()[-600:] or '未返回成功的PPTX文件'))
            receipt = record
        exported_path = inside(self.folder, receipt['output'])
        report = json.loads((self.folder / 'validation' / (exported_path.stem + '.report.json')).read_text(encoding='utf-8'))
        if (report.get('schema') != 'ppt-master.pptx-postflight-report.v1'
                or report.get('status') not in {'passed', 'passed-with-warnings'}
                or Path((report.get('output') or {}).get('path', '')).resolve() != exported_path
                or (report.get('output') or {}).get('bytes') != exported_path.stat().st_size
                or (report.get('source') or {}).get('svg_slide_count') != len(names)):
            raise ValueError('实际导出文件与原版导出报告不一致')
        deck = Presentation(exported_path)
        if len(deck.slides) != len(names):
            raise ValueError('实际PPTX页数与原版SVG页面队列不一致')
        # This is a product index, not a required model-authored artifact.
        manifest_file = self.folder / 'page_manifest.json'
        try:
            previous_metadata = {page['file']: page for page in json.loads(manifest_file.read_text(encoding='utf-8'))['pages']}
        except (OSError, ValueError, KeyError, TypeError):
            previous_metadata = {}
        pages = [dict(previous_metadata.get(name) or {}, file=name) for name in names]
        manifest = {'pages': pages}
        if not 3 <= len(pages) <= 20:
            raise ValueError('当前工作区支持3—20页演示文稿')
        expected = (state['brief'].get('page_count') if session.get('revision_action') == 'rebuild' else
                    state.get('effective_page_count') or state['brief'].get('page_count'))
        if session.get('revision_action') in {None, 'rebuild'} and expected and len(pages) != expected:
            raise ValueError('PPT Master页面清单与要求的总页数不一致')
        if session.get('revision_action'):
            previous = session['previous_pages']
            selected = set(session['target_ids'])
            action = session['revision_action']
            snapshot = session.get('previous_snapshot') or {}
            previous_root = Path(session['previous_project']) if session.get('previous_project') else None

            def before_bytes(page):
                """原地修订用修订前快照比对，不依赖旧工程副本。"""
                if previous_root is not None:
                    return (previous_root / 'svg_output' / page['master_file']).read_bytes()
                return snapshot_bytes(snapshot, page['master_file'])

            if action == 'replace' and names != [p['master_file'] for p in previous]:
                raise ValueError('替换操作必须保留原页面数量、顺序和文件标识')
            if action == 'split' and len(pages) <= len(previous):
                raise ValueError('拆分操作必须增加页面')
            if action == 'append' and len(pages) <= len(previous):
                raise ValueError('追加操作必须增加页面')
            old_deck = (Presentation(Path(previous_root).parent / 'presentation.pptx')
                        if previous_root is not None else None)
            if names == [page['master_file'] for page in previous] and all(
                    (self.folder / 'svg_output' / page['master_file']).read_bytes() == before_bytes(page)
                    and (old_deck is None or
                         page_content(deck.slides[index]) == page_content(old_deck.slides[index]))
                    for index, page in enumerate(previous)):
                raise ValueError('尚未实际修改页面或讲稿，未交付复制的旧版本')
            for page in previous:
                if page['slide_id'] not in selected:
                    if (self.folder / 'svg_output' / page['master_file']).read_bytes() != before_bytes(page):
                        raise ValueError('修订改变了未选页面，未交付新版本')
                    if old_deck is None:
                        continue
                    index = names.index(page['master_file'])
                    old_index = previous.index(page)
                    if page_content(deck.slides[index]) != page_content(old_deck.slides[old_index]):
                        raise ValueError('实际导出改变了未选页内容，未交付新版本')
        evidence = state.get('evidence') or {}
        analyses = {a['analysis_id']: a for a in state.get('analyses') or []}
        materials = {m['material_id']: m for m in state.get('materials') or []}
        old_ids = {p['master_file']: p['slide_id'] for p in session.get('previous_pages', [])}
        slides = []
        for index, (page, native) in enumerate(zip(pages, deck.slides)):
            notes = native.notes_slide.notes_text_frame.text if native.has_notes_slide else ''
            ids = list(dict.fromkeys([*(page.get('source_ids') or []), *(eid for eid in evidence if eid in notes)]))
            page['source_ids'] = ids
            if any(eid not in evidence for eid in ids):
                raise ValueError('页面来源不属于当前任务')
            texts = [shape.text for shape in native.shapes if shape.has_text_frame and shape.text.strip()]
            charts = [shape.chart for shape in native.shapes if shape.has_chart]
            tables = [shape.table for shape in native.shapes if shape.has_table]
            analysis = analyses.get(page.get('analysis_id'))
            if page.get('analysis_id') and not analysis:
                raise ValueError('页面引用了不存在的分析记录')
            if charts and analysis:
                if len(charts) != 1 or len(charts[0].series) != 1:
                    raise ValueError('已绑定分析的图表结构不一致')
                if analysis['material_id'] not in ids:
                    raise ValueError('图表页须引用确定性分析的原始材料')
                if [Decimal(str(v)) for v in charts[0].series[0].values] != [Decimal(v) for v in analysis['values']]:
                    raise ValueError('实际图表数值与确定性分析不一致')
                if [str(c.label) for c in charts[0].plots[0].categories] != analysis['categories']:
                    raise ValueError('实际图表分类与确定性分析不一致')
            if page.get('analysis_id') and not charts:
                raise ValueError('要求的原生图表尚未导出为可编辑数据图表')
            material = materials.get(page.get('table_material_id'))
            if page.get('table_material_id') and not material:
                raise ValueError('表格引用的数据材料不属于当前任务')
            if material:
                if material['material_id'] not in ids:
                    raise ValueError('表格页须引用原始数据材料')
                actual = [[[cell.text for cell in row.cells] for row in table.rows] for table in tables]
                wanted = [[str(v) if v is not None else '' for v in row]
                          for row in [material['columns'], *material['rows']]]
                if wanted not in actual:
                    raise ValueError('原生表格与原始数据不一致')
            notes_enabled = '--no-notes' not in receipt['arguments']
            if notes_enabled and ids and not native.has_notes_slide:
                raise ValueError('页面来源须写入PPT备注')
            if notes_enabled and any(eid not in notes for eid in ids):
                raise ValueError('PPT备注缺少页面来源ID')
            from uuid import uuid4
            sid = old_ids.get(page['file']) or page.get('slide_id') or uuid4().hex[:12]
            if not isinstance(sid, str) or not re.fullmatch(r'[A-Za-z0-9_-]{1,64}', sid):
                raise ValueError('页面稳定标识格式无效')
            slides.append({'slide_id': sid, 'master_file': page['file'], 'kind': 'master',
                           'title': page.get('title') or (texts[0] if texts else '第%d页' % (index + 1)),
                           'purpose': page.get('purpose', ''), 'claims': [],
                           'speaker_notes': notes.strip(),
                           'notes': [{'text': '本页引用资料' if notes_enabled else '资料面板可查看本页引用',
                                      'evidence_ids': ids}] if ids else []})
        if len({s['slide_id'] for s in slides}) != len(slides):
            raise ValueError('修订页面标识重复')
        for page, slide in zip(pages, slides):
            page['slide_id'] = slide['slide_id']
        manifest_text = json.dumps(manifest, ensure_ascii=False, indent=2)
        if not manifest_file.is_file() or manifest_file.read_text(encoding='utf-8') != manifest_text:
            manifest_file.write_text(manifest_text, encoding='utf-8')
        output = self.folder.parent / 'presentation.pptx'
        shutil.copyfile(exported_path, output)
        previews = render_previews(output, self.folder.parent / 'previews', [p['slide_id'] for p in slides])
        return {'slides': slides, 'revision': session['revision'], 'effective_page_count': len(slides),
                'stale_slide_ids': [], 'pending_user': None,
                'artifacts': {'pptx_path': str(output), 'pages': len(slides), 'previews': previews,
                              'revision': session['revision'], 'renderer': 'ppt-master'}}
