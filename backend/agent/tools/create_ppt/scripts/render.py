"""将已核验的内容渲染为 PPTX；模型不参与文件操作。"""
from pathlib import Path
import os
import json

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

from .qa import check


INK = RGBColor(27, 42, 61)
ACCENT = RGBColor(26, 107, 117)
FONT = "Microsoft YaHei" if os.name == "nt" else "Noto Sans CJK SC"


def sources_markdown(evidence: dict, slides) -> str:
    """仅从结构化台账建立来源清单，不反解析 observation。"""
    ids = {eid for slide in slides for claim in [*slide.claims, *slide.notes]
           for eid in claim.evidence_ids}
    ids.update(data["material_id"] for slide in slides for data in (slide.chart, slide.table)
               if data and data.get("material_id"))
    lines = []
    for eid in sorted(ids & evidence.keys()):
        item = evidence[eid]
        title = item.get("title") or "未命名来源"
        heading = " / ".join(item.get("heading_path") or [])
        line = item.get("line")
        loc = " / ".join(p for p in (heading, "第%d行" % line if line else "") if p)
        excerpt = " ".join((item.get("text") or "").split())[:180]
        lines.append("- [%s] %s%s%s" % (eid, title, " · " + loc if loc else "",
                                        " · " + excerpt if excerpt else ""))
    return "# 资料来源\n\n" + "\n".join(lines) + "\n"


def _textbox(slide, left, top, width, height, text, size=21, bold=False, color=INK):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.text = text
    for paragraph in frame.paragraphs:
        paragraph.font.size = Pt(size)
        paragraph.font.name = FONT
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
    return box


def render(slides, path, evidence: dict, sources_md: str | None = None) -> Path:
    """QA 不通过则不写文件；调用方负责选择安全的目标路径。"""
    sources_md = sources_md if sources_md is not None else sources_markdown(evidence, slides)
    report = check(slides, evidence, sources_md)
    if not report.ok:
        raise ValueError("PPT QA 未通过：%s" % [(x.code, x.slide) for x in report.issues])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    for index, item in enumerate(slides, 1):
        page = prs.slides.add_slide(blank)
        _textbox(page, 0.7, 0.45, 11.9, 0.85, item.title, 32, True, INK)
        if item.kind in ("cover", "end"):
            if item.subtitle:
                _textbox(page, 0.75, 2.5, 11.5, 1.6, item.subtitle, 24, False, ACCENT)
        elif item.kind == "agenda":
            for i, text in enumerate(item.items):
                _textbox(page, 1.0, 1.45 + i * 0.6, 11.3, 0.5,
                         "%02d  %s" % (i + 1, text), 21)
        elif item.kind == "bullets":
            for i, claim in enumerate(item.claims):
                _textbox(page, 0.95, 1.35 + i * 0.83, 11.5, 0.73,
                         "•  " + claim.text, 20)
        elif item.kind == "compare":
            _textbox(page, 0.85, 1.3, 5.7, 0.5, item.left_title, 22, True, ACCENT)
            _textbox(page, 6.85, 1.3, 5.7, 0.5, item.right_title, 22, True, ACCENT)
            for side, left in (("left", 0.85), ("right", 6.85)):
                group = [c for c in item.claims if c.column == side]
                for i, claim in enumerate(group):
                    _textbox(page, left, 2.0 + i * 1.05, 5.6, 0.9,
                             "•  " + claim.text, 18)
        elif item.kind == "chart":
            data = item.chart
            chart_data = CategoryChartData()
            chart_data.categories = data["categories"]
            chart_data.add_series(data.get("label") or "数值", [float(v) for v in data["values"]])
            chart = page.shapes.add_chart(
                XL_CHART_TYPE.COLUMN_CLUSTERED if data.get("chart_type") == "bar" else XL_CHART_TYPE.LINE,
                Inches(0.85), Inches(1.6), Inches(11.7), Inches(3.65), chart_data).chart
            chart.has_legend = False
            chart.has_title = False
            chart.chart_style = 10
            if data.get("chart_type") == "bar" and min(float(v) for v in data["values"]) >= 0:
                chart.value_axis.minimum_scale = 0
            for axis in (chart.category_axis, chart.value_axis):
                axis.tick_labels.font.name = FONT
                axis.tick_labels.font.size = Pt(15)
            chart.value_axis.has_title = bool(data.get("unit"))
            if data.get("unit"):
                chart.value_axis.axis_title.text_frame.text = data["unit"]
            for i, claim in enumerate(item.claims):
                _textbox(page, 0.95, 5.4 + i * 0.36, 11.5, 0.38, claim.text, 17)
        elif item.kind == "table":
            data = item.table
            units = data.get("units") or {}
            columns = [c + "（" + units[c] + "）" if units.get(c) and units[c] not in c else c for c in data["columns"]]
            rows = [columns, *data["rows"]]
            table = page.shapes.add_table(len(rows), len(rows[0]), Inches(0.8), Inches(1.55),
                                          Inches(11.7), Inches(0.48 * len(rows))).table
            for ri, row in enumerate(rows):
                for ci, value in enumerate(row):
                    cell = table.cell(ri, ci)
                    cell.text = "" if value is None else str(value)
                    for paragraph in cell.text_frame.paragraphs:
                        paragraph.font.name = FONT
                        paragraph.font.size = Pt(17)
                        paragraph.font.bold = ri == 0
                    if ri == 0:
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = ACCENT
                        for paragraph in cell.text_frame.paragraphs:
                            paragraph.font.color.rgb = RGBColor(255, 255, 255)
            for i, claim in enumerate(item.claims):
                _textbox(page, 0.95, 6.0 + i * 0.4, 11.5, 0.4, claim.text, 17)
        _textbox(page, 11.7, 7.0, 0.7, 0.25, str(index), 10, False, ACCENT)
        cited = {eid for claim in [*item.claims, *item.notes] for eid in claim.evidence_ids}
        cited.update(data["material_id"] for data in (item.chart, item.table) if data)
        source_notes = ["来源：%s；%s%s%s%s" % (evidence[eid].get("title", ""),
                        " / ".join(evidence[eid].get("heading_path") or []),
                        "；" + evidence[eid]["url"] if evidence[eid].get("url") else "",
                        "；适用范围：" + evidence[eid]["scope"] if evidence[eid].get("scope") else "",
                        "；日期：" + evidence[eid]["date"] if evidence[eid].get("date") else "")
                        for eid in sorted(cited) if eid in evidence]
        if item.chart:
            source_notes.append("数据口径：" + json.dumps(item.chart, ensure_ascii=False))
        if source_notes:
            _textbox(page, 0.75, 7.0, 10.7, 0.28,
                     "来源：" + "；".join(dict.fromkeys(evidence[eid].get("title", "") for eid in sorted(cited) if eid in evidence)),
                     10, False, ACCENT)
        page.notes_slide.notes_text_frame.text = "\n".join([*(c.text for c in item.notes), *source_notes])
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return output
