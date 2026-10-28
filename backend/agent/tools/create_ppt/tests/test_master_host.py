"""Hosted workflow contracts; fake model responses never make paid requests."""
import copy
import json
from pathlib import Path
import sys

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[4]))
from core.config import settings
from core.errors import ApiError
from agent.tools.create_ppt.scripts import delivery, master_host, master_workspace, workflow
from agent.tools.create_ppt.scripts.materials import Material


def reply(name, **arguments):
    return {'text': '', 'calls': [{'name': name, 'arguments': arguments}],
            'usage': {'input_tokens': 1, 'output_tokens': 1}}


@pytest.fixture
def state(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, 'agent_ppt_workdir', str(tmp_path))
    def initialize(workspace, state):
        workspace.folder.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(master_workspace.MasterWorkspace, 'initialize', initialize)
    return workflow.create({'goal': '说明课程示例', 'audience': '新生', 'page_count': 3},
                           [Material(title='材料', text='只含两门课程示例').model_dump()], 1, None)


def test_new_tasks_and_default_generation_use_master(state, monkeypatch):
    assert state['renderer'] == 'ppt-master'
    called = []
    monkeypatch.setattr(master_host, 'advance', lambda current: called.append(current) or current)
    workflow.generate(state['task_id'], 1, 0)
    assert len(called) == 1 and called[0]['renderer'] == 'ppt-master'


def test_two_confirmations_require_actual_answers_and_stale_answer_does_not_reserve(state):
    first = master_host.advance(state, complete=lambda *a, **k: reply('request_confirmation',
        stage='stage1', question='确认目标？', proposal='面向新生，只讲材料中的两门课程'))
    assert first['stage'] == 'await_design' and not first['master']['approved']
    token = first['pending_user']['confirmation_id']
    with pytest.raises(ApiError):
        workflow.continue_master(state['task_id'], 1, 0,
                                 {'confirmation_id': 'old', 'approved': True, 'feedback': ''})
    assert delivery.owned_state(state['task_id'], 1)['stage'] == 'await_design'
    second = master_host.advance(first, confirmation={'confirmation_id': token, 'approved': True},
        complete=lambda *a, **k: reply('request_confirmation', stage='stage2', question='确认设计？', proposal='三页方案'))
    assert second['master']['approved'] == ['stage1']
    assert second['pending_user']['stage'] == 'stage2'
    pending = copy.deepcopy(second['pending_user'])
    with pytest.raises(ValueError, match='确认'):
        master_host._execute(second, 'write_file', {'path': 'svg_output/01.svg', 'text': '<svg />'})
    second = master_host.advance(second, confirmation={'confirmation_id': pending['confirmation_id'],
        'approved': False, 'feedback': '请减少文字'}, complete=lambda *a, **k: reply('request_confirmation',
        stage='stage2', question='确认调整方案？', proposal='减少文字后的方案'))
    assert second['master']['approved'] == ['stage1']


def test_cannot_skip_confirmations_or_forge_upstream_receipts(state):
    master_host.start(state)
    with pytest.raises(ValueError, match='Stage 1'):
        master_host._execute(state, 'request_confirmation', {'stage': 'stage2', 'question': '确认', 'proposal': '方案'})
    with pytest.raises(ValueError, match='伪造'):
        master_host._execute(state, 'write_file', {'path': 'result.json', 'text': '{}'})
    with pytest.raises(ValueError, match='尚未'):
        master_host._execute(state, 'finish', {})


def test_model_steps_are_persisted_without_replaying_completed_tools(state, monkeypatch):
    executed = []
    original = master_host._execute
    def execute(current, name, arguments):
        executed.append(arguments.get('path', name))
        return original(current, name, arguments)
    monkeypatch.setattr(master_host, '_execute', execute)
    response = lambda *a, **k: reply('list_files', path='')
    first = master_host.advance(state, complete=response)
    assert first['stage'] == 'master_continue' and len(executed) == 2
    restored = delivery.owned_state(state['task_id'], 1)
    second = master_host.advance(restored, complete=lambda *a, **k: reply('request_confirmation',
        stage='stage1', question='确认？', proposal='方案'))
    assert len(executed) == 4  # one remaining list, then one confirmation
    assert second['stage'] == 'await_design' and second['master']['api_requests'] == 4


def test_failed_script_returns_evidence_not_success(state, monkeypatch):
    master_host.start(state)
    state['master']['approved'] = ['stage1', 'stage2']
    monkeypatch.setattr(master_workspace, 'command', lambda *a: {'exit_code': 2, 'output': 'blocked'})
    result = master_host._execute(state, 'run_script', {'script': 'svg_to_pptx.py', 'arguments': ['@project']})
    assert result['exit_code'] == 2
    with pytest.raises(ValueError, match='导出成功'):
        master_host._execute(state, 'finish', {})


def test_workspace_paths_and_script_arguments_cannot_escape(tmp_path):
    workspace = master_workspace.MasterWorkspace(tmp_path)
    with pytest.raises(ValueError):
        workspace.read('../.env')
    with pytest.raises(ValueError):
        workspace.write('../outside.md', 'x')
    with pytest.raises(ValueError):
        workspace.write('svg_output/page.svg', '<svg><image href="file:///root/.bashrc"/></svg>')
    with pytest.raises(ValueError):
        workspace.run('svg_to_pptx.py', ['@project', '-o', '/root/out.pptx'])
    with pytest.raises(ValueError):
        workspace.run('anything.py', [])


def test_selected_revision_rejects_writing_other_pages(state):
    master_host.start(state)
    state['master'].update(approved=['stage1', 'stage2'], revision_action='replace', target_ids=['p02'],
        previous_pages=[{'slide_id': 'p01', 'master_file': '01.svg'}, {'slide_id': 'p02', 'master_file': '02.svg'}])
    with pytest.raises(ValueError, match='未选'):
        master_host._execute(state, 'write_file', {'path': 'svg_output/01.svg', 'text': '<svg />'})
    assert master_host._execute(state, 'write_file', {'path': 'svg_output/02.svg', 'text': '<svg />'})['written']


def test_actual_upstream_initialization_writes_materials_and_analysis(tmp_path):
    workspace = master_workspace.MasterWorkspace(tmp_path / 'deck_20261025')
    material = Material(title='课程', columns=['课程', '学时'], rows=[['操作系统', 40]], units={'学时': '学时'})
    workspace.initialize({'brief': {'goal': '课程示例'}, 'materials': [material.model_dump()],
                          'evidence': {material.material_id: material.evidence()}, 'analyses': [{'values': ['40']}]})
    assert (workspace.folder / 'svg_output').is_dir()
    assert '操作系统' in (workspace.folder / 'sources' / (material.material_id + '.csv')).read_text(encoding='utf-8')
    assert json.loads((workspace.folder / 'sources/analyses.json').read_text()) == [{'values': ['40']}]


def test_chat_handoff_uses_master_and_emits_existing_task_event(state, monkeypatch):
    import asyncio
    from agent.tools.create_ppt import tool
    from llm import llm_api_client
    monkeypatch.setattr(llm_api_client, 'complete_sync', lambda *a, **k: reply('request_confirmation',
        stage='stage1', question='确认目标？', proposal='校园材料讲解'))
    tid = delivery.reserve_task(1, 902)
    async def collect():
        return [event async for event in tool.run_stream([], '制作PPT',
            task_id=tid, owner_id=1, conversation_id=902,
            brief={'goal': '讲解校园材料', 'audience': '新生'})]
    events = asyncio.run(collect())
    event = next(payload['ppt_task'] for kind, payload in events if kind == 'ppt')
    assert event['renderer'] == 'ppt-master' and event['stage'] == 'await_design'
    assert sum(kind == 'done' for kind, _ in events) == 1
    assert delivery.owned_state(tid, 1)['brief']['audience'] == '新生'


def test_changed_materials_invalidate_pending_design(state):
    waiting = master_host.advance(state, complete=lambda *a, **k: reply('request_confirmation',
        stage='stage1', question='确认？', proposal='方案'))
    token = waiting['pending_user']['confirmation_id']
    workflow.add_materials(state['task_id'], 1, [Material(title='补充材料', text='新的适用范围')])
    with pytest.raises(ApiError):
        workflow.continue_master(state['task_id'], 1, 0, {'confirmation_id': token, 'approved': True})


def test_publication_checks_real_file_preview_and_revision_token(state, monkeypatch):
    """Native-file fixture checks delivery only, not model/upstream authoring quality."""
    from pptx import Presentation
    from pptx.util import Inches
    from agent.tools.create_ppt.scripts import preview
    master_host.start(state)
    folder = Path(state['master']['project'])
    (folder / 'svg_output').mkdir(exist_ok=True)
    (folder / 'exports').mkdir(exist_ok=True)
    source_id = next(iter(state['evidence']))
    pages = []
    deck = Presentation()
    for n in range(3):
        name = '%02d_page.svg' % n
        (folder / 'svg_output' / name).write_text('<svg />')
        pages.append({'file': name, 'title': '页面', 'purpose': '说明材料', 'source_ids': [source_id]})
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = '两门课程示例'
        slide.notes_slide.notes_text_frame.text = source_id
    deck.save(folder / 'exports/result.pptx')
    (folder / 'page_manifest.json').write_text(json.dumps({'pages': pages}))
    observed = []
    def previews(path, output, ids):
        assert len(Presentation(path).slides) == 3
        observed.append(path)
        output.mkdir(exist_ok=True)
        files = []
        for sid in ids:
            file = output / (sid + '.png')
            file.write_bytes(b'test-preview')
            files.append(str(file))
        return files
    monkeypatch.setattr(preview, 'render_previews', previews)
    result = master_workspace.MasterWorkspace(folder).publish(state, state['master'])
    assert observed and result['artifacts']['renderer'] == 'ppt-master'
    assert len({s['slide_id'] for s in result['slides']}) == 3
    state.update(result, stage='draft')
    master_host._save(state)
    _, token = delivery.confirm(state['task_id'], 1, expected_revision=1)
    state = delivery.owned_state(state['task_id'], 1)
    state.update(stage='draft', revision=2)
    master_host._save(state)
    with pytest.raises(ApiError):
        delivery.consume(token, 1)
    # Upstream commands could change an exported page even when its SVG stayed
    # byte-identical: the publication check must compare the actual native page.
    import shutil
    changed = folder.parent.parent / 'v2' / 'project_changed_20261025'
    shutil.copytree(folder, changed)
    altered = Presentation(changed / 'exports/result.pptx')
    altered.slides[0].shapes[0].text = '不应改变的页面被修改'
    altered.save(changed / 'exports/result.pptx')
    session = {**state['master'], 'revision': 2, 'revision_action': 'replace',
               'previous_project': str(folder), 'previous_pages': result['slides'],
               'target_ids': [result['slides'][1]['slide_id']]}
    with pytest.raises(ValueError, match='实际导出改变'):
        master_workspace.MasterWorkspace(changed).publish(state, session)


def test_replaced_page_retains_old_project_for_pause_and_resume(state, monkeypatch):
    from llm import llm_api_client
    master_host.start(state)
    state.update(stage='draft', revision=1, slides=[{'slide_id': 'p01', 'master_file': '01.svg'}],
                 artifacts={'pptx_path': 'unused'})
    master_host._save(state)
    previous = state['master']['project']
    monkeypatch.setattr(llm_api_client, 'complete_sync', lambda *a, **k: reply('list_files', path=''))
    updated = workflow.revise(state['task_id'], 1, 1, '只修改第一张', ['p01'])
    assert updated['stage'] == 'master_continue' and updated['revision'] == 1
    assert updated['master']['previous_project'] == previous
    assert updated['master']['target_ids'] == ['p01']


def test_sync_transport_sends_skill_tools_and_parses_function_calls(monkeypatch):
    import httpx
    from llm import llm_api_client
    captured = []
    original = httpx.Client
    def transport(request):
        captured.append(json.loads(request.content))
        return httpx.Response(200, json={'status': 'completed', 'id': 'local-fixture',
            'model': 'fixture', 'usage': {}, 'output': [{'type': 'function_call',
            'call_id': 'call_fixture', 'name': 'list_files', 'arguments': '{"path":""}'}]})
    config = llm_api_client.get_api_config()
    monkeypatch.setenv(config.active.api_key_env, 'unit-test-only')
    monkeypatch.setattr(httpx, 'Client', lambda **kwargs: original(transport=httpx.MockTransport(transport), **kwargs))
    result = llm_api_client.complete_sync([{'role': 'user', 'content': '继续制作'}],
                                          task='ppt', tools=master_host.TOOLS)
    assert captured[0]['tools'][0]['type'] == 'function'
    assert captured[0]['max_output_tokens'] == config.tasks['ppt'].max_output_tokens
    assert result['calls'] == [{'name': 'list_files', 'arguments': {'path': ''}}]


def test_upstream_batch_helpers_receive_stdin(tmp_path, monkeypatch):
    received = []
    monkeypatch.setattr(master_workspace, 'command', lambda *args:
        received.append(args) or {'exit_code': 0, 'output': '<g />'})
    workspace = master_workspace.MasterWorkspace(tmp_path)
    result = workspace.run('preset_shape_svg.py', ['render-batch', '--input', '-'], '{"shapes": []}')
    assert result['exit_code'] == 0 and received[0][-1] == '{"shapes": []}'


def test_campus_lookup_reuses_registry_and_preserves_task_evidence(state, monkeypatch):
    from agent.runtime import registry
    from agent.core.contracts import Source, ToolResult
    master_host.start(state)
    monkeypatch.setattr(registry, 'call', lambda *args: ToolResult(tool='search_regulations', ok=True,
        sources=[Source(doc_id='policy', title='本校规定', text='2025级适用新规则。')]))
    arguments = {'tool': 'search_regulations', 'arguments': {'query': '适用年级'}, 'question': '哪些年级适用？'}
    with pytest.raises(ValueError, match='材料范围'):
        master_host._execute(state, 'query_campus', arguments)
    state['brief']['material_policy'] = 'campus'
    result = master_host._execute(state, 'query_campus', arguments)
    assert result['status'] == 'success' and result['sources'][0]['text'] == '2025级适用新规则。'
    assert state['budget']['research_steps_left'] == 7


def test_failed_model_attempt_still_consumes_budget_and_can_pause(state):
    def failed(*a, **k):
        raise RuntimeError('offline transport failure')
    with pytest.raises(RuntimeError):
        master_host.advance(state, complete=failed)
    restored = delivery.owned_state(state['task_id'], 1)
    assert restored['stage'] == 'paused' and restored['master']['api_requests'] == 1
