"""阶段4：离线内容、渲染与 QA；坏样本必须让检查器报错。"""
import sys
from pathlib import Path

import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[4]))

from agent.tools.create_ppt.scripts.content.model import Claim, Slide  # noqa: E402
from agent.tools.create_ppt.scripts.rendering.qa import check  # noqa: E402
from agent.tools.create_ppt.scripts.rendering.render import render, sources_markdown  # noqa: E402


EID = "0123456789ab"
EVIDENCE = {EID: {"title": "某培养方案", "heading_path": ["教学计划"],
                  "line": 42, "text": "课程学分为2"}}


def _good():
    return [Slide("cover", "课程定位", subtitle="面向教师"),
            Slide("agenda", "目录", items=["课程信息", "对比"]),
            Slide("bullets", "课程信息", [Claim("课程学分为2", [EID])],
                  notes=[Claim("依据培养方案第42行", [EID])]),
            Slide("compare", "专业对比",
                  [Claim("甲专业课程学分为2", [EID], column="left"),
                   Claim("乙专业暂无依据", gap="本次未检索到", column="right")],
                  left_title="甲专业", right_title="乙专业"),
            Slide("end", "谢谢")]


def test_known_bad_input_must_be_visible_to_qa():
    bad = [Slide("bullets", "坏样本", [Claim("无来源的事实")])]
    assert "missing_source" in {x.code for x in check(bad, {}, "").issues}


def test_roundtrip_and_claim_level_sources(tmp_path):
    slides = _good()
    md = sources_markdown(EVIDENCE, slides)
    assert EID in md and "某培养方案" in md
    assert check(slides, EVIDENCE, md).ok
    path = render(slides, tmp_path / "demo.pptx", EVIDENCE, md)
    prs = Presentation(path)
    assert len(prs.slides) == 5
    texts = [shape.text for page in prs.slides for shape in page.shapes if shape.has_text_frame]
    for expected in ("课程定位", "课程学分为2", "甲专业课程学分为2", "乙专业暂无依据"):
        assert any(expected in text for text in texts)
    assert "依据培养方案第42行" in prs.slides[2].notes_slide.notes_text_frame.text


@pytest.mark.parametrize("slides,evidence,md,code", [
    ([Slide("bullets", "过长", [Claim("字" * 106, [EID])])], EVIDENCE,
     "- [0123456789ab] 来源\n", "overflow"),
    ([Slide("bullets", "空")], {}, "", "empty_page"),
    ([Slide("bullets", "标题", [Claim("事实")])], {}, "", "missing_source"),
    ([Slide("bullets", "标题", [Claim("事实", [EID])])], {},
     "- [0123456789ab] 来源\n", "dangling_evidence"),
    ([Slide("cover", "封面")], EVIDENCE, "- [0123456789ab] 来源\n", "unused_source"),
    ([Slide("bullets", "标题", [Claim("事实", [EID])])], EVIDENCE,
     "", "unlisted_evidence"),
    ([Slide("compare", "标题", [Claim("事实", [EID], column="left")] * 5,
            left_title="甲", right_title="乙")], EVIDENCE,
     "- [0123456789ab] 来源\n", "overflow"),
])
def test_bad_samples_are_detected(slides, evidence, md, code):
    assert code in {x.code for x in check(slides, evidence, md).issues}


def test_claim_free_pages_exempt_and_invalid_render_writes_nothing(tmp_path):
    assert check([Slide("cover", "封面"), Slide("end", "结束")], {}, "").ok
    path = tmp_path / "bad.pptx"
    with pytest.raises(ValueError, match="QA 未通过"):
        render([Slide("bullets", "标题", [Claim("无来源")])], path, {})
    assert not path.exists()
