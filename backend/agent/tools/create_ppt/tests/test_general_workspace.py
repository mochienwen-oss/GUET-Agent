"""Focused offline checks for the general workspace; no model or network calls."""
import io
import json
import copy
from pathlib import Path
import sys
from concurrent.futures import ThreadPoolExecutor
from threading import Event

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[4]))
from core.config import settings
from core.errors import ApiError
from agent.runtime import loop
from agent.tools.create_ppt.scripts import delivery, master_host, master_workspace, subgraph, workflow
from agent.tools.create_ppt.scripts.materials import Material, read_materials


@pytest.fixture(autouse=True)
def offline(tmp_path, monkeypatch):
    import socket

    def forbidden(*args, **kwargs):
        raise AssertionError('This offline check must not access the network or a real API')

    monkeypatch.setattr(socket.socket, 'connect', forbidden)
    monkeypatch.setattr(loop, 'complete_sync', forbidden)
    monkeypatch.setattr(settings, 'agent_ppt_workdir', str(tmp_path))

    def initialize(workspace, state):
        workspace.folder.mkdir(parents=True, exist_ok=True)
        workspace.import_materials(state)

    monkeypatch.setattr(master_workspace.MasterWorkspace, 'initialize', initialize)


def task(policy='general', materials=None):
    return workflow.create({'goal': '介绍时间管理', 'audience': '一般听众', 'material_policy': policy},
                           materials or [], 1)


def response(text='', name=None, **arguments):
    calls = [{'id': 'offline-call', 'name': name, 'arguments': arguments}] if name else []
    message = {'role': 'assistant', 'content': text}
    if name:
        message['tool_calls'] = [{'id': 'offline-call', 'type': 'function', 'function': {
            'name': name, 'arguments': json.dumps(arguments)}}]
    return {'text': text, 'calls': calls, 'message': message, 'usage': {}}


@pytest.mark.parametrize('policy', ['general', 'public', 'provided_only', 'campus'])
def test_sources_are_available_without_frontend_mode_routing(policy, monkeypatch):
    material = Material(title='参考内容', text='先确定目标，再安排时间').model_dump()
    state = task(policy, [material] if policy == 'provided_only' else [])
    available = []

    def model(messages, **kwargs):
        available.extend(item['function']['name'] for item in kwargs['tools'])
        return response('请确认主题方向与内容结构。')

    monkeypatch.setattr(loop, 'complete_sync', model)
    result = workflow.generate(state['task_id'], 1, 0)
    assert result['stage'] == 'await_design'
    assert {'search_regulations', 'search_web', 'read_webpage'} <= set(available)
    assert {'read_file', 'write_file', 'edit_file', 'execute_python'} <= set(available)


def test_material_only_mode_cannot_start_without_materials():
    state = task('provided_only')
    assert workflow.generate(state['task_id'], 1, 0)['stage'] == 'await_materials'


def test_real_confirmation_and_revision_feedback_are_preserved(monkeypatch):
    state = task()
    monkeypatch.setattr(loop, 'complete_sync', lambda *a, **k: response('大纲与设计方案'))
    first = workflow.generate(state['task_id'], 1, 0)
    token = first['pending_user']['confirmation_id']
    with pytest.raises(ApiError):
        workflow.continue_master(state['task_id'], 1, 0,
                                 {'confirmation_id': 'stale', 'approved': True})
    assert delivery.owned_state(state['task_id'], 1)['stage'] == 'await_design'
    second = workflow.continue_master(state['task_id'], 1, 0,
        {'confirmation_id': token, 'approved': True, 'feedback': '我编辑后的方案：先讲目标，再讲方法'})
    assert second['master']['approved'] == ['stage1']
    assert second['pending_user']['stage'] == 'stage2'
    files = master_host._files(second)
    assert files.execute('write_file', {'path': 'svg_output/01.svg', 'text': '<svg/>'})['status'] == 'error'
    third = workflow.continue_master(state['task_id'], 1, 0, {
        'confirmation_id': second['pending_user']['confirmation_id'], 'approved': True, 'feedback': ''})
    assert third['master']['approved'] == ['stage1', 'stage2']
    assert any('先讲目标，再讲方法' in item.get('content', '') for item in third['master']['messages'])
    assert master_host._files(third).execute('write_file', {
        'path': 'svg_output/01.svg', 'text': '<svg/>'})['status'] == 'success'


def test_failed_file_reads_pause_automatic_continuation(monkeypatch):
    state = task()
    monkeypatch.setattr(loop, 'complete_sync', lambda *a, **k: response(name='read_file', path='missing.md'))
    result = workflow.generate(state['task_id'], 1, 0)
    assert result['stage'] == 'paused'
    assert result['master']['api_requests'] == 3
    assert '文件不存在' in result['errors'][-1]['detail']


def test_uploaded_image_is_available_only_in_its_task(tmp_path):
    from PIL import Image
    state = task()
    uploads = tmp_path / state['task_id'] / 'uploads'
    stream = io.BytesIO()
    Image.new('RGBA', (12, 8), (10, 20, 30, 128)).save(stream, format='PNG')
    materials = read_materials('配图.png', stream.getvalue(), asset_dir=uploads)
    state = workflow.add_materials(state['task_id'], 1, materials)
    master_host.start(state)
    project = Path(state['master']['project'])
    assert (project / 'images' / materials[0].asset_file).is_file()
    assert 'images/' in (project / 'sources' / (materials[0].material_id + '.md')).read_text(encoding='utf-8')
    files = master_host._files(state)
    outside = tmp_path / 'other-user.txt'
    outside.write_text('private', encoding='utf-8')
    assert files.execute('read_file', {'path': str(outside)})['status'] == 'error'
    assert 'sources/' in files.read_file('@project')['data']['entries']
    with pytest.raises(ValueError, match='图片须'):
        read_materials('配图.png', stream.getvalue())


def test_same_task_cannot_run_twice_concurrently(monkeypatch):
    state = task()
    entered, release = Event(), Event()

    def wait(current):
        entered.set()
        assert release.wait(5)
        return current

    monkeypatch.setattr(master_host, 'advance', wait)
    with ThreadPoolExecutor(max_workers=1) as executor:
        first = executor.submit(workflow.generate, state['task_id'], 1, 0)
        try:
            assert entered.wait(5)
            with pytest.raises(ApiError):
                workflow.generate(state['task_id'], 1, 0)
        finally:
            release.set()
        first.result()
    assert state['task_id'] not in workflow._ACTIVE_TASKS


def test_interrupted_task_is_recoverable_after_restart():
    state = task()
    master_host.start(state)
    state['stage'] = 'generating'
    subgraph.save_state(state['task_id'], state)
    subgraph.recover_interrupted()
    restored = delivery.owned_state(state['task_id'], 1)
    assert restored['stage'] == 'paused'
    assert restored['next_node'] == 'master'
    assert '不自动重放' in restored['master']['messages'][-1]['content']


@pytest.mark.skipif(sys.platform != 'win32', reason='Windows sharing violation')
def test_short_windows_file_lock_does_not_lose_state(monkeypatch):
    state = task()
    replace = subgraph.os.replace
    attempts = []

    def sharing(source, destination):
        attempts.append(destination)
        if len(attempts) < 3:
            error = PermissionError('sharing violation')
            error.winerror = 32
            raise error
        return replace(source, destination)

    monkeypatch.setattr(subgraph.os, 'replace', sharing)
    monkeypatch.setattr(subgraph.time, 'sleep', lambda delay: None)
    state['activity'] = '已保存的步骤'
    subgraph.save_state(state['task_id'], state)
    assert len(attempts) == 3
    assert delivery.owned_state(state['task_id'], 1)['activity'] == '已保存的步骤'


def test_chat_creation_and_first_message_are_separate_without_duplicate(monkeypatch):
    state = workflow.create_chat('', 1)
    assert state['chat_messages'] == []
    monkeypatch.setattr(master_host, 'advance', lambda current, **kwargs: current)
    result = workflow.message(state['task_id'], 1, 0, '做6页时间管理汇报')
    assert [entry['content'] for entry in result['chat_messages']] == ['做6页时间管理汇报']


def test_active_save_does_not_restore_a_cleared_or_deleted_selection():
    state = workflow.create_chat('', 1)
    state['slides'] = [{'slide_id': 'page3', 'title': '案例'}]
    state['selected_slide_id'] = 'page3'
    subgraph.save_state(state['task_id'], state, selection_update=True)
    old_snapshot = copy.deepcopy(state)
    workflow.select_slide(state['task_id'], 1, None)
    subgraph.save_state(state['task_id'], old_snapshot)
    assert delivery.owned_state(state['task_id'], 1)['selected_slide_id'] is None
    workflow.select_slide(state['task_id'], 1, 'page3')
    old_snapshot['slides'] = []
    subgraph.save_state(state['task_id'], old_snapshot)
    assert delivery.owned_state(state['task_id'], 1)['selected_slide_id'] is None


def test_messages_and_continuations_preserve_cumulative_usage(monkeypatch):
    state = task()
    master_host.start(state)
    state['master'].update(api_requests=60, usage=[{'input_tokens': 3}])
    state['stage'] = 'failed'
    subgraph.save_state(state['task_id'], state)

    def continue_without_model(current, **kwargs):
        current['stage'] = 'master_continue'
        subgraph.save_state(current['task_id'], current)
        return current

    monkeypatch.setattr(master_host, 'advance', continue_without_model)
    result = workflow.message(state['task_id'], 1, 0, '请根据已保存的内容继续')
    assert result['master']['api_requests'] == 60
    assert result['master']['usage'] == [{'input_tokens': 3}]
    result['master']['api_requests'] = 7
    subgraph.save_state(result['task_id'], result)
    resumed = workflow.continue_master(result['task_id'], 1, 0)
    assert resumed['master']['api_requests'] == 7


def test_user_pause_waits_for_worker_and_keeps_returned_usage(monkeypatch):
    import asyncio
    from agent.tools.create_ppt.scripts import background

    state = task()
    entered, release = Event(), Event()

    def in_flight_model(*args, **kwargs):
        entered.set()
        if not release.wait(5):
            raise TimeoutError('Offline response was not released')
        result = response(name='write_file', path='outline.md', text='Must not be written after pause')
        result['usage'] = {'input_tokens': 11, 'output_tokens': 2}
        return result

    monkeypatch.setattr(loop, 'complete_sync', in_flight_model)

    async def scenario():
        background.start(workflow.message, state['task_id'], 1, 0, '制作一份课堂汇报')
        worker = background._RUNNING[state['task_id']]
        try:
            assert await asyncio.to_thread(entered.wait, 3)
            background.pause(state['task_id'], 1, 0)
            assert background.running(state['task_id']) and background.pausing(state['task_id'])
            assert state['task_id'] in workflow._ACTIVE_TASKS
            with pytest.raises(ApiError):
                background.start(workflow.continue_master, state['task_id'], 1, 0, manual=True)

            # Pausing one task must not cancel a separate task's scope.
            other = task()
            def complete_other(task_id, owner_id, revision):
                current = delivery.owned_state(task_id, owner_id)
                current['stage'] = 'await_design'
                subgraph.save_state(task_id, current)
                return current
            background.start(complete_other, other['task_id'], 1, 0)
            await asyncio.wait_for(asyncio.shield(background._RUNNING[other['task_id']]), 3)
            assert delivery.owned_state(other['task_id'], 1)['stage'] == 'await_design'
            assert background.running(state['task_id'])

            release.set()
            await asyncio.wait_for(asyncio.shield(worker), 3)
            saved = delivery.owned_state(state['task_id'], 1)
            assert saved['stage'] == 'paused' and saved['pause_requested']
            assert workflow.can_continue(saved)
            assert saved['master']['api_requests'] == 1
            assert saved['master']['usage'] == [{'input_tokens': 11, 'output_tokens': 2}]
            assert not (Path(saved['master']['project']) / 'outline.md').exists()
            assert not saved.get('errors')
            assert saved['chat_messages'][-1]['content'].startswith('已暂停制作')
            pending = saved['master']['messages'][-1]
            assert pending['role'] == 'tool' and json.loads(pending['content'])['error']['code'] == 'interrupted'
            assert not background.running(state['task_id'])
            assert state['task_id'] not in workflow._ACTIVE_TASKS
        finally:
            release.set()
            await background.shutdown()

    asyncio.run(scenario())


def test_pause_before_first_step_does_not_start_model_and_is_resumable():
    import asyncio
    from agent.tools.create_ppt.scripts import background

    state = task()
    async def scenario():
        def must_not_run(*args):
            pytest.fail('A task paused before starting must not run its operation')
        background.start(must_not_run, state['task_id'], 1, 0)
        worker = background._RUNNING[state['task_id']]
        background.pause(state['task_id'], 1, 0)
        await asyncio.wait_for(asyncio.shield(worker), 3)
        saved = delivery.owned_state(state['task_id'], 1)
        assert saved['stage'] == 'paused' and workflow.can_continue(saved)
        assert not saved.get('master')
    asyncio.run(scenario())


def test_pause_checks_owner_and_revision_before_touching_a_task():
    from agent.tools.create_ppt.scripts import background
    state = task()
    with pytest.raises(ApiError):
        background.pause(state['task_id'], 2, 0)
    with pytest.raises(ApiError):
        background.pause(state['task_id'], 1, 1)
    assert background.pause(state['task_id'], 1, 0)['stage'] == state['stage']
    assert not background.pausing(state['task_id'])


def test_task_cancel_event_survives_normal_execution_segments():
    from agent.core.request_control import check_cancelled, request_scope
    stop = Event()
    with request_scope(5, cancel_event=stop):
        check_cancelled()
    assert not stop.is_set()
    with request_scope(5, cancel_event=stop):
        stop.set()
        with pytest.raises(InterruptedError):
            check_cancelled()
    check_cancelled()


def test_independent_chat_can_be_created_while_another_task_is_unfinished():
    first = workflow.create_chat('', 1)
    first['stage'] = 'generating'
    subgraph.save_state(first['task_id'], first)
    second = workflow.create_chat('', 1)
    assert second['task_id'] != first['task_id']


def test_material_read_by_active_host_does_not_enter_upload_path(monkeypatch):
    state = task()
    master_host.start(state)
    state['stage'] = 'generating'

    def forbidden(*args, **kwargs):
        raise AssertionError('Active host must not use the ordinary upload mutation')

    monkeypatch.setattr(workflow, 'add_materials', forbidden)
    material = Material(title='公开原始资料', text='实际读取的正文', source_type='public_web', url='https://example.org/source')
    receipt = master_host._register_material(state, material, {
        'final_url': material.url, 'accessed_at': '2026-10-26T12:00:00Z',
        'published_at': None, 'content_scope': 'webpage', 'reading_status': 'read'})
    assert receipt['evidence_id'] == material.material_id
    assert state['evidence'][material.material_id]['published_at'] is None
    assert state['evidence'][material.material_id]['accessed_at'] == '2026-10-26T12:00:00Z'


def test_discussion_is_not_confirmation_but_real_nonce_button_is(monkeypatch):
    state = workflow.create_chat('', 1)
    monkeypatch.setattr(loop, 'complete_sync', lambda *args, **kwargs: response('请确认当前内容方向'))
    first = workflow.message(state['task_id'], 1, 0, '制作课堂汇报')
    second = workflow.message(state['task_id'], 1, 0, '同意，继续',
                              confirmation_id=first['pending_user']['confirmation_id'])
    assert second['master']['approved'] == []
    assert second.get('confirmation_records', []) == []
    assert delivery.owned_state(state['task_id'], 1)['chat_messages'] == second['chat_messages']
    nonce = second['pending_user']['confirmation_id']
    count = len(second['chat_messages'])
    with pytest.raises(ApiError):
        workflow.message(state['task_id'], 1, 0, '', confirmation_id='stale', approved=True)
    assert len(delivery.owned_state(state['task_id'], 1)['chat_messages']) == count
    third = workflow.message(state['task_id'], 1, 0, '', confirmation_id=nonce, approved=True)
    assert third['master']['approved'] == ['stage1']
    assert third['confirmation_records'][-1]['confirmation_id'] == nonce
    assert third['confirmation_records'][-1]['approved'] is True


def saved_deck():
    """Small editable fixture; no Skill, provider or desktop renderer involved."""
    from pptx import Presentation
    state = task()
    master_host.start(state)
    project = Path(state['master']['project'])
    (project / 'svg_output').mkdir()
    deck, slides = Presentation(), []
    for number in range(1, 4):
        filename = '%02d.svg' % number
        (project / 'svg_output' / filename).write_text('<svg><text>page%d</text></svg>' % number, encoding='utf-8')
        native = deck.slides.add_slide(deck.slide_layouts[6])
        native.shapes.add_textbox(0, 0, 914400, 914400).text = 'page%d' % number
        slides.append({'slide_id': 'p%d' % number, 'master_file': filename, 'title': 'page%d' % number})
    path = project.parent / 'presentation.pptx'
    deck.save(path)
    state.update(revision=1, effective_page_count=3, slides=slides, stage='draft',
                 artifacts={'pptx_path': str(path), 'pages': 3, 'previews': []},
                 selected_slide_id='p2', revision_pending=True,
                 last_user_message='按照我的本轮要求修改作品')
    state['master']['approved'] = ['stage1', 'stage2']
    subgraph.save_state(state['task_id'], state, selection_update=True)
    return state


def test_whole_deck_scope_is_not_restricted_by_selected_page():
    state = saved_deck()
    old_project = Path(state['master']['project'])
    receipt = master_host._set_revision_scope(state, 'rebuild', ['p2'], '用户要求整份重做')
    assert receipt['data']['slide_ids'] == ['p1', 'p2', 'p3']
    assert state['master']['approved'] == []
    assert list((Path(state['master']['project']) / 'svg_output').glob('*.svg')) == []
    assert len(list((old_project / 'svg_output').glob('*.svg'))) == 3


def test_rebuild_checks_current_page_target_instead_of_old_effective_count():
    state = saved_deck()
    master_host._set_revision_scope(state, 'rebuild', [], '用户要求重做为4页')
    state['brief']['page_count'] = 4
    project = Path(state['master']['project'])
    (project / 'page_manifest.json').write_text(json.dumps({'pages': [
        {'file': '%02d.svg' % number, 'title': 'page%d' % number, 'source_ids': []}
        for number in range(1, 4)]}), encoding='utf-8')
    with pytest.raises(ValueError, match='总页数'):
        master_workspace.MasterWorkspace(project).publish(state, state['master'])


def test_revision_receipt_follows_real_export_and_preview_files(monkeypatch):
    from PIL import Image
    from pptx import Presentation
    from agent.tools.create_ppt.scripts import preview
    state = saved_deck()
    old_file = Path(state['artifacts']['pptx_path'])
    master_host._set_revision_scope(state, 'replace', ['p2'], '用户要求仅修改第二页')
    project = Path(state['master']['project'])
    assert master_host._files(state).execute('write_file', {'path': 'svg_output/01.svg', 'text': '<svg/>'})['status'] == 'error'
    (project / 'svg_output' / '02.svg').write_text('<svg><text>changed</text></svg>', encoding='utf-8')
    (project / 'page_manifest.json').write_text(json.dumps({'pages': [
        {'file': '%02d.svg' % number, 'title': 'page%d' % number, 'source_ids': []}
        for number in range(1, 4)]}), encoding='utf-8')
    deck = Presentation(old_file)
    deck.slides[1].shapes[0].text = 'changed'
    (project / 'exports').mkdir()
    deck.save(project / 'exports' / 'fixture.pptx')
    state['master']['commands'] = [{'script': 'svg_to_pptx.py', 'exit_code': 0}]

    def fixture_previews(pptx, output, ids):
        output.mkdir(parents=True)
        for sid in ids:
            Image.new('RGB', (2, 2), 'white').save(output / (sid + '.png'))
        return [str(output / (sid + '.png')) for sid in ids]

    monkeypatch.setattr(preview, 'render_previews', fixture_previews)
    monkeypatch.setattr(loop, 'complete_sync', lambda *args, **kwargs: response('模型的完成说明'))
    result = master_host.advance(state)
    assert result['stage'] == 'draft' and result['revision'] == 2
    assert Path(result['artifacts']['pptx_path']).is_file()
    assert all(Path(path).is_file() for path in result['artifacts']['previews'])
    assert [page['slide_id'] for page in result['chat_messages'][-1]['updated_pages']] == ['p2']
    assert '第2页' in result['chat_messages'][-1]['content']
    assert Presentation(old_file).slides[1].shapes[0].text == 'page2'


def test_model_claim_cannot_publish_without_an_actual_export(monkeypatch):
    state = saved_deck()
    master_host._set_revision_scope(state, 'replace', ['p2'], '用户要求修改第二页')
    project = Path(state['master']['project'])
    (project / 'page_manifest.json').write_text(json.dumps({'pages': [
        {'file': '%02d.svg' % number, 'title': 'page%d' % number, 'source_ids': []}
        for number in range(1, 4)]}), encoding='utf-8')
    state['master']['commands'] = [{'script': 'svg_to_pptx.py', 'exit_code': 0}]
    monkeypatch.setattr(loop, 'complete_sync', lambda *args, **kwargs: response('修改完成'))
    with pytest.raises(ValueError, match='实际导出'):
        master_host.advance(state)
    assert state['stage'] == 'paused' and state['revision'] == 1
    assert '修改完成' not in state['chat_messages'][-1]['content']


def test_reading_new_content_continues_but_covered_read_only_round_pauses(monkeypatch):
    state = task()
    monkeypatch.setattr(loop, 'complete_sync', lambda *args, **kwargs: response(
        name='read_file', path='@project/sources/brief.json', line_count=5))
    first = workflow.generate(state['task_id'], 1, 0)
    assert first['stage'] == 'master_continue'
    assert first['master']['api_requests'] == 3
    second = workflow.continue_master(state['task_id'], 1, 0)
    assert second['stage'] == 'paused'
    assert second['master']['api_requests'] == 6
    assert second['errors'][-1]['node'] == 'reading'
    assert '重复阅读未形成新方案' in second['chat_messages'][-1]['content']
    monkeypatch.setattr(loop, 'complete_sync', lambda *args, **kwargs: response(
        name='read_file', path='@project/sources/brief.json', start_line=5, line_count=5))
    resumed = workflow.continue_master(state['task_id'], 1, 0)
    assert resumed['stage'] == 'master_continue'
    assert resumed['master']['api_requests'] == 9
