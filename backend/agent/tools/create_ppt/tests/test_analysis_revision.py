"""Real PPTX objects and task state; generation/preview doubles are explicit."""
import copy
import sys
from dataclasses import replace
from pathlib import Path

import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[4]))

from core.config import settings
from core.errors import ApiError
from agent.tools.create_ppt.scripts import delivery, workflow
from agent.tools.create_ppt.scripts.materials import Material
from agent.tools.create_ppt.scripts.model import Claim, Slide


@pytest.fixture
def task(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "agent_ppt_workdir", str(tmp_path))
    # Deliberate preview double: these tests validate lifecycle, not visual quality.
    def preview(path, output, ids):
        output.mkdir(parents=True, exist_ok=True)
        from PIL import Image
        paths = []
        for sid in ids:
            target = output / (sid + ".png")
            Image.new("RGB", (4, 4), "white").save(target)
            paths.append(str(target))
        return paths
    monkeypatch.setattr(workflow, "render_previews", preview)
    a = Material(title="旧年级规定", text="2024级不执行退学预警", scope="2024级")
    b = Material(title="新年级规定", text="2025级累计25学分退学预警", scope="2025级")
    data = Material(title="课程学时", columns=["课程", "学时"], rows=[["操作系统", 40], ["计算机网络", 48]], source_type="data")
    state = workflow.create({"goal": "新生说明", "audience": "2025级本科生", "page_count": 5},
                            [m.model_dump() for m in (a, b, data)], 1, 123)
    analysis = workflow.add_analysis(state["task_id"], 1, data.material_id,
                                    {"operation": "sum", "field": "学时", "group_by": "课程", "unit": "学时"})
    def author(state):
        return [Slide("cover", "新生说明", slide_id="cover"),
                Slide("compare", "区分年级", [Claim(a.text, [a.material_id], column="left"),
                      Claim(b.text, [b.material_id], column="right")], left_title="2024级", right_title="2025级", slide_id="cohorts"),
                Slide("chart", "课程学时", chart=analysis, slide_id="hours"),
                Slide("table", "原始学时", table={"material_id": data.material_id, "columns": data.columns, "rows": data.rows}, slide_id="raw"),
                Slide("end", "按适用年级办理", slide_id="end")]
    def no_research(brief):
        return {"research_questions": [], "research_tasks": [], "evidence_needs": []}
    state = workflow.generate(state["task_id"], 1, 0, planner=no_research, authorer=author,
                              tool_call=lambda *a: pytest.fail("Provided sources must not trigger retrieval"))
    assert state["stage"] == "draft", state.get("errors")
    return state


def test_mixed_materials_generate_native_chart_table_and_sources(task):
    deck = Presentation(task["artifacts"]["pptx_path"])
    assert len(deck.slides) == 5 and task["revision"] == 1
    chart = next(shape.chart for shape in deck.slides[2].shapes if shape.has_chart)
    assert list(chart.series[0].values) == [40.0, 48.0]
    assert chart.value_axis.minimum_scale == 0
    table = next(shape.table for shape in deck.slides[3].shapes if shape.has_table)
    assert table.cell(1, 0).text == "操作系统" and table.cell(1, 1).text == "40"
    assert "2024级" in deck.slides[1].notes_slide.notes_text_frame.text
    assert len(task["artifacts"]["previews"]) == 5
    assert not task["tool_calls"]


def test_targeted_revision_preserves_other_pages_and_invalidates_old_token(task):
    tid = task["task_id"]
    _, old = delivery.confirm(tid, 1, expected_revision=1)
    original = copy.deepcopy(task["slides"])
    updated = workflow.revise(tid, 1, 1, "第四页标题改为原始数据", ["raw"],
          authorer=lambda state, ids, feedback: [replace(workflow.read_slide(state["slides"][3]), title="原始数据")])
    assert updated["revision"] == 2 and updated["stage"] == "draft"
    for index in (0, 1, 2, 4):
        assert updated["slides"][index] == original[index]
    assert updated["slides"][3]["slide_id"] == "raw" and updated["slides"][3]["title"] == "原始数据"
    with pytest.raises(ApiError):
        delivery.consume(old, 1)
    with pytest.raises(ApiError):
        delivery.confirm(tid, 1, expected_revision=1)
    _, token = delivery.confirm(tid, 1, expected_revision=2)
    assert delivery.consume(token, 1) == Path(updated["artifacts"]["pptx_path"])


def test_split_and_append_keep_stable_ids_and_brief(task):
    tid = task["task_id"]
    original = copy.deepcopy(task)
    split = workflow.revise(tid, 1, 1, "第三页拆两页", ["hours"], action="split",
        authorer=lambda state, ids, feedback: [workflow.read_slide(state["slides"][2]),
            Slide("bullets", "阅读数据", [Claim("建议先比较总学时", kind="suggestion")])])
    assert len(split["slides"]) == 6 and split["slides"][2]["slide_id"] == "hours"
    assert split["slides"][3]["slide_id"] not in {s["slide_id"] for s in original["slides"]}
    assert split["slides"][4] == original["slides"][3] and split["brief"] == original["brief"]
    appended = workflow.revise(tid, 1, 2, "前面不变，补行动建议", [], action="append",
        authorer=lambda *a: [Slide("bullets", "行动建议", [Claim("建议核对个人所属年级", kind="suggestion")])])
    assert appended["slides"][:-1] == split["slides"]


def test_material_change_requires_updating_related_pages(task):
    old = task["materials"][1]
    state = workflow.add_materials(task["task_id"], 1,
        [Material(title=old["title"], text="2025级按实际细则执行")], replace_id=old["material_id"])
    assert state["stale_slide_ids"] == ["cohorts"]
    with pytest.raises(ValueError):
        workflow.revise(task["task_id"], 1, 1, "只改第四页", ["raw"], authorer=lambda *a: [])
    with pytest.raises(ApiError):
        delivery.confirm(task["task_id"], 1, expected_revision=1)


def test_failed_revision_keeps_previous_version_and_rejects_wrong_owner(task):
    tid = task["task_id"]
    with pytest.raises(ValueError):
        workflow.revise(tid, 1, 1, "错误替换", ["raw"], authorer=lambda *a: [])
    current = workflow._owned(tid, 1)
    assert current["revision"] == 1 and current["slides"] == task["slides"]
    with pytest.raises(ApiError):
        workflow.revise(tid, 2, 1, "修改", ["raw"], authorer=lambda *a: [])


def test_stale_revision_is_rejected_before_model_request(task, monkeypatch):
    from agent.tools.create_ppt.scripts import generate
    monkeypatch.setattr(generate, "_json_reply", lambda *a: pytest.fail("Stale request must not call the model"))
    with pytest.raises(ApiError):
        workflow.revise(task["task_id"], 1, 0, "第三页简化")


def test_brief_change_invalidates_confirmation_and_requires_full_generation(task):
    state = workflow.update_brief(task["task_id"], 1, 1, {**task["brief"], "audience": "学院教师"})
    assert state["stage"] == "brief_changed" and state["brief"]["audience"] == "学院教师"
    with pytest.raises(ApiError):
        delivery.confirm(task["task_id"], 1, expected_revision=1)
