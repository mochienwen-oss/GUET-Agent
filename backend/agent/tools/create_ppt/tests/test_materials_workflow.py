"""Material/data contracts; no model or network requests."""
import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[4]))

from agent.tools.create_ppt.scripts.materials import Material, read_materials
from agent.tools.create_ppt.scripts.analysis import analyse


def test_two_documents_keep_distinct_scope_and_locations():
    from docx import Document
    doc = Document()
    doc.add_paragraph("2024级学生不执行退学预警条款。")
    stream = io.BytesIO()
    doc.save(stream)
    first = read_materials("2024说明.docx", stream.getvalue())[0]
    second = read_materials("2025说明.md", "2025级学生按新细则执行。".encode())[0]
    assert first.material_id != second.material_id
    assert "2024" in first.text and "2025" not in first.text
    assert "2025" in second.text and second.locator


def test_csv_and_xlsx_preserve_headers_and_numeric_cells():
    from openpyxl import Workbook
    csv = read_materials("学时.csv", "课程,理论,实验\n操作系统,32,8\n计算机网络,32,16\n".encode())[0]
    wb = Workbook()
    wb.active.append(["课程", "理论", "实验"])
    wb.active.append(["操作系统", 32, 8])
    stream = io.BytesIO()
    wb.save(stream)
    xlsx = read_materials("学时.xlsx", stream.getvalue())[0]
    assert csv.columns == xlsx.columns == ["课程", "理论", "实验"]
    assert analyse(csv, {"operation": "sum", "field": "实验", "unit": "学时"})["values"] == ["24"]
    assert analyse(xlsx, {"operation": "sum", "field": "理论"})["values"] == ["32"]


def test_numeric_analysis_preserves_filter_group_unit_and_denominator():
    data = Material(title="调查", columns=["组", "人数", "总人数"],
                    rows=[["甲", 12, 20], ["乙", 6, 10]], source_type="data")
    result = analyse(data, {"operation": "share", "field": "人数", "denominator_field": "总人数",
                            "denominator_aggregation": "sum", "group_by": "组", "filters": {"组": "甲"}, "unit": "%"})
    assert result["values"] == ["60"] and result["categories"] == ["甲"]
    assert result["denominators"] == ["20"] and result["row_count"] == 1
    assert result["material_id"] == data.material_id and result["filters"] == {"组": "甲"}
    assert analyse(data, {"operation": "mean", "field": "人数"})["values"] == ["9"]
    assert analyse(data, {"operation": "count"})["values"] == ["2"]


@pytest.mark.parametrize("rows,spec", [
    ([["甲", None]], {"operation": "sum", "field": "人数"}),
    ([["甲", "未知"]], {"operation": "mean", "field": "人数"}),
    ([["甲", 12]], {"operation": "share", "field": "人数"}),
])
def test_missing_number_or_denominator_is_not_guessed(rows, spec):
    with pytest.raises(ValueError):
        analyse(Material(title="统计", columns=["组", "人数"], rows=rows), spec)


def test_unsupported_or_empty_document_is_not_silently_accepted():
    with pytest.raises(ValueError):
        read_materials("材料.exe", b"x")
    with pytest.raises(ValueError):
        read_materials("空.txt", b"")


def test_sufficient_sources_allow_zero_research_and_policy_cannot_be_overridden(monkeypatch):
    from agent.tools.create_ppt.scripts import generate
    monkeypatch.setattr(generate, "_json_reply", lambda *a: {"research_tasks": [], "gaps": []})
    assert generate.plan({"goal": "材料说明", "material_policy": "provided_only"})["research_tasks"] == []
    monkeypatch.setattr(generate, "_json_reply", lambda *a: {"research_tasks": [
        {"question": "额外查询", "name": "search_regulations", "args": {"query": "额外查询"}}]})
    with pytest.raises(ValueError, match="仅用"):
        generate.plan({"goal": "材料说明", "material_policy": "provided_only"})


def test_share_requires_explicit_denominator_aggregation():
    data = Material(title="同一总人数重复填写", columns=["组", "人数", "总人数"], rows=[["甲", 10, 20], ["乙", 10, 20]])
    spec = {"operation": "share", "field": "人数", "denominator_field": "总人数"}
    with pytest.raises(ValueError, match="分母.*口径"):
        analyse(data, spec)
    assert analyse(data, {**spec, "denominator_aggregation": "single"})["values"] == ["100"]
    assert analyse(data, {**spec, "denominator_aggregation": "sum"})["values"] == ["50"]
