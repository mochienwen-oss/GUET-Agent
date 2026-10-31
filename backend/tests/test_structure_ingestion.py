"""Necessary structure, score and upload-path checks; no exact indicator queries."""
import asyncio
import io
from pathlib import Path
from types import SimpleNamespace

import pytest

from documents.catalog import document_chunks
from documents.ingest import mdify
from documents.ingest import pdfextract
from agent.tools.search_regulations.scripts.index_ops import indexing
from agent.tools.search_regulations.scripts.recall import retrieval
from agent.tools.search_regulations.scripts.evidence_ops import gate_signals
from agent.tools.search_regulations.scripts.evidence_ops.evidence_gate import POLICY_V1, decide
from agent.tools.search_regulations.scripts.recall.rerank_scores import (
    SCORE_VERSION, sigmoid, threshold_on_single_scale,
)


TABLE = """<table><tr><th>课程名称</th><th>学分</th><th>备注</th></tr>
<tr><td rowspan="2">计算机网络</td><td>3</td><td>必修</td></tr>
<tr><td>3</td><td></td></tr>
<tr><td colspan="3">同一说明不得重复三次</td></tr></table>"""
DOC = "# 样本文档\n\n## 课程设置\n\n适用范围及表外条件完整保留。\n\n" + TABLE


def test_unique_cells_keep_spans_blanks_and_external_conditions():
    rows = [u for u in document_chunks.build_chunks(DOC) if u["block_type"] == "table"]
    assert len(rows) == 3
    assert rows[-1]["text"].count("同一说明不得重复三次") == 1
    assert len(rows[-1]["table_fields"]) == 1
    assert rows[-1]["table_fields"][0]["columns"] == [1, 2, 3]
    assert rows[1]["table_fields"][0]["inherited_from_rowspan"]
    assert "〔空白，原表未填写〕" in rows[1]["text"]
    assert all("表外条件完整保留" in row["text"] for row in rows)
    assert all("<td" not in row["text"] for row in rows)
    assert all(DOC[slice(*row["source_span"]) ] == TABLE for row in rows)


def test_no_size_boundary_or_silent_article_truncation():
    article = "### 第五条\n" + "条件与后果必须保留。" * 300
    units = indexing._build_index_units(article)
    assert len(units) == 1 and units[0]["text"] == article
    assert indexing.split_text(article, max_len=30) == [article]


def test_regulation_mentioning_graduation_is_not_misclassified_as_plan():
    text = "## 三、课程认定办法\n\n依据课程目标和毕业要求认定。\n\n（一）必备条件及其后果不可拆开。"
    assert len(document_chunks.build_chunks(text, "规章制度")) == 1
    assert len(document_chunks.build_chunks(text)) == 1


def test_invalid_span_is_rejected_not_guessed():
    with pytest.raises(ValueError):
        document_chunks.build_chunks('<table><tr><td rowspan="3">不完整</td></tr></table>')


def test_unknown_header_keeps_first_row():
    units = document_chunks.build_chunks('<table><tr><td>甲</td><td>乙</td></tr></table>')
    assert len(units) == 1 and "甲" in units[0]["text"]
    assert units[0]["review_flags"]


def test_pdf_geometry_and_layout_line_breaks_survive_normalization():
    rows = [["课程名称", "备注"], ["计算机网络\n操作系统", "完整注释"], [None, "另行说明"]]
    boxes = [[(0, 0, 1, 1), (1, 0, 2, 1)], [(0, 1, 1, 3), (1, 1, 2, 2)],
             [None, (1, 2, 2, 3)]]
    table = pdfextract._table_to_md(rows, boxes)
    assert 'rowspan="2"' in table and "计算机网络<br>操作系统" in table
    normalized = mdify._plain_to_md("第一条 借用手续\n\n" + table)
    assert table in normalized and "### 第一条" in normalized
    chunks = document_chunks.build_chunks(normalized)
    assert "计算机网络\n操作系统" in next(u for u in chunks if u["block_type"] == "table")["text"]


def test_html_upload_keeps_table_geometry():
    md = mdify.to_markdown("test.html", ("<h1>课程</h1>"+TABLE).encode())
    assert 'rowspan="2"' in md and 'colspan="3"' in md
    assert len([u for u in document_chunks.build_chunks(md) if u["block_type"] == "table"]) == 3


def test_docx_merge_retains_real_geometry_and_line_breaks():
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "课程名称"
    table.cell(0, 1).text = "学分"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "计算机网络\n操作系统"
    table.cell(1, 1).text = "3"
    table.cell(2, 1).text = "4"
    raw = mdify._docx_table_to_md(table)
    assert 'rowspan="2"' in raw and "计算机网络<br>操作系统" in raw
    units = document_chunks.build_chunks(raw)
    assert len(units) == 2
    assert units[1]["table_fields"][0]["inherited_from_rowspan"]


def test_scanned_pdf_does_not_masquerade_as_complete_markdown(monkeypatch):
    monkeypatch.setattr(pdfextract, "to_markdown_report", lambda *a, **kw:
                        ("只有页眉的残缺文字", {"unprocessed_scan_pages": [1, 2]}))
    with pytest.raises(ValueError, match="未完成OCR"):
        mdify.to_markdown("扫描.pdf", b"fake-pdf")


def test_rerank_uses_identity_then_exactly_one_sigmoid(monkeypatch):
    calls = []
    class Model:
        tokenizer = staticmethod(lambda *args, **kw: {"input_ids": [1, 2]})
        max_length = 8192
        def predict(self, pairs, **kwargs):
            from torch.nn import Identity
            assert isinstance(kwargs["activation_fn"], Identity)
            calls.append(pairs)
            return [2.0]
    monkeypatch.setattr(retrieval, "_get_reranker", lambda: Model())
    rows = [{"source": "样本", "snippet": "计算机网络：3学分"}]
    retrieval._rerank("课程多少学分？", rows)
    assert len(calls) == 1
    assert rows[0]["rerank_score"] == pytest.approx(sigmoid(2.0))
    assert rows[0]["rerank_logit"] == 2.0
    assert rows[0]["rerank_score_version"] == SCORE_VERSION


def test_score_unit_conversion_preserves_gate_without_recalibration():
    assert threshold_on_single_scale(.54) == pytest.approx(.16034265007517948)
    assert threshold_on_single_scale(.70) == pytest.approx(.8472978603872034)
    for logit in (-8, -2, 0, 1, 2, 5):
        score = sigmoid(logit)
        old = {"text": "原文", "rerank_score": sigmoid(score)}
        new = {"text": "原文", "rerank_score": score, "rerank_score_version": SCORE_VERSION}
        first = decide(gate_signals.collect(query="课程", ranked=[old]), POLICY_V1)
        second = decide(gate_signals.collect(query="课程", ranked=[new]), POLICY_V1)
        assert (first.status, first.reason) == (second.status, second.reason)


def test_oversized_rerank_is_not_silently_truncated(monkeypatch):
    model = SimpleNamespace(max_length=2, tokenizer=lambda *a, **kw: {"input_ids": [1, 2, 3]})
    monkeypatch.setattr(retrieval, "_get_reranker", lambda: model)
    with pytest.raises(ValueError, match="禁止静默截断"):
        retrieval._rerank("普通问题", [{"snippet": "完整结构"}])


def test_admin_upload_preview_and_create_share_production_splitter(monkeypatch):
    from fastapi import BackgroundTasks, UploadFile
    from routers import admin
    from routers.schemas import DocumentCreate
    parsed = asyncio.run(admin.parse_document(UploadFile(filename="测试.md", file=io.BytesIO(DOC.encode())), user=None))
    assert parsed["chunking_version"] == document_chunks.VERSION
    body = DocumentCreate(title="测试", content=parsed["markdown"])
    preview = admin.preview_document(body, user=None)
    assert preview["chunks"] == indexing._build_index_units(DOC)
    background = BackgroundTasks()
    response = admin.create_document(body, background, user=None)
    assert response["status"] == "processing" and len(background.tasks) == 1
    # Do not execute the task against the actual MySQL/vector store in this test.


def test_http_upload_and_edited_preview_contract(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from routers import admin
    app = FastAPI()
    app.include_router(admin.router, prefix="/admin")
    app.dependency_overrides[admin.require_admin] = lambda: None
    submitted = []
    monkeypatch.setattr(admin, "_ingest_document", lambda *args: submitted.append(args))
    with TestClient(app) as client:
        result = client.post("/admin/documents/parse", files={"file": ("课程.md", DOC.encode(), "text/markdown")})
        assert result.status_code == 200
        parsed = result.json()
        assert parsed["table_rows"] == 3 and parsed["warnings"] == []
        body = {"title": "课程", "content": parsed["markdown"], "category": "培养方案"}
        preview = client.post("/admin/documents/preview", json=body)
        assert preview.status_code == 200
        assert preview.json()["chunks"] == indexing._build_index_units(DOC, "培养方案")
        saved = client.post("/admin/documents", json=body)
        assert saved.status_code == 200 and saved.json()["status"] == "processing"
        assert len(submitted) == 1 and submitted[0][1] == DOC


def test_production_llm_messages_receive_fields_not_html_source():
    from agent.core import ToolResult
    from agent.runtime import registry, graph
    row = next(u for u in indexing._build_index_units(DOC) if u["block_type"] == "table")
    context = retrieval._chunk_from_point(SimpleNamespace(payload={**row, "source": "课程方案"}), .8)
    observation = registry.render(ToolResult(tool="search_regulations", ok=True, status="success",
                                            data={"blocks": [context]}))
    calls = [{"name": "search_regulations", "arguments": {"query": "计算机网络多少学分？"}}]
    messages = graph.init_messages([], "计算机网络多少学分？") + graph.round_messages(calls, [observation], 0)
    assert messages[-1]["role"] == "tool" and messages[-1]["tool_call_id"] == "call_0"
    assert row["text"].strip() in messages[-1]["content"]
    assert "课程名称：计算机网络" in observation and "学分：3" in observation
    assert "<td" not in observation and "<table" not in observation


def test_sync_never_turns_table_projection_into_original_markdown(monkeypatch):
    from sqlalchemy import create_engine, Text
    from sqlalchemy.orm import sessionmaker
    from database import db
    engine = create_engine("sqlite://")
    # Only the in-memory test dialect changes; production keeps MySQL LONGTEXT.
    monkeypatch.setattr(db.Document.__table__.c.content, "type", Text())
    db.Document.__table__.create(engine)
    sessions = sessionmaker(bind=engine)
    units = indexing._build_index_units(DOC)
    points = [SimpleNamespace(payload={**u, "document_id": 42, "source": "课程方案"}) for u in units]
    monkeypatch.setattr(indexing, "ensure_collection", lambda: None)
    monkeypatch.setattr(indexing, "_get_client", lambda: SimpleNamespace(scroll=lambda **kw: (points, None)))
    monkeypatch.setattr(db, "SessionLocal", sessions)
    assert indexing.sync_documents()["created"] == 1
    with sessions() as session:
        assert session.get(db.Document, 42).content == ""
    engine.dispose()


def test_all_reviewed_source_documents_have_reconstructible_spans():
    root = Path(__file__).resolve().parents[2] / "final-md"
    paths = list(root.rglob("*.md"))
    assert len(paths) == 92
    for path in paths:
        text = path.read_text(encoding="utf-8")
        units = document_chunks.build_chunks(text, "培养方案" if path.parent.name == "培养方案" else "规章制度")
        assert units, path
        for unit in units:
            raw = text[slice(*unit["source_span"])]
            if unit["block_type"] == "paragraph":
                assert raw == unit["text"]
            else:
                assert raw.startswith(("<table", "|"))
                assert "<td" not in unit["text"]
